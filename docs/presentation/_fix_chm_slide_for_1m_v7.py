"""Slide 21 said the corn rows are missing data. At 1 m they are gone.

THE PROBLEM THIS FIXES
----------------------
Slide 21 read:

    The corn rows are NOT canopy.
    They are missing data - 3.2% of this window, where no first return came back.

That was true of the 0.5 m CHM it used to show. Every terrain figure has now
been rebuilt at 1 m, where the same window has 0.00% no-data and no stripes at
all, so the slide's own picture contradicted its text.

THE FIX
-------
Show both, using `chm_window_multires_9t.png` -- the same 300 m window built at
0.5 m, 1 m and 2 m side by side, with the no-data cells picked out and the
percentage under each panel. The slide stops asserting the artefact and starts
demonstrating it, which is stronger, and the claim now matches what is on
screen.

The figure is 2.54:1, so the old layout (a full-height picture down the right
third) cannot hold it. Text moves to a wide block at the top and the figure runs
the full width underneath.

Numbers come from the figure itself, so they cannot drift apart:
  0.5 m  3.23% of cells have no value
  1 m    0.00%
  2 m    0.00%
These are interior-only, which is why they differ slightly from the whole-tile
CHM void of 2.04% -> 0.03%.

Colour: the figure is a greyscale canopy ramp with orange no-data. No red, no
green, so the CLAUDE.md pair rule cannot be violated, and the no-data cells
carry a legend entry as well as a colour.

Run:
    python docs/presentation/_fix_chm_slide_for_1m_v7.py
Writes (in place):
    docs/presentation/WellSight_Presentation v7.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs/presentation/WellSight_Presentation v7.pptx"
FIG = (ROOT / "docs/presentation/figures_30to45min/v6"
       / "chm_window_multires_9t.png")

EXPECT = "Step 1: terrain derivatives"
SLIDE = 21

TITLE = "Step 1: terrain derivatives – canopy height"
BULLETS = [
    "How tall are the green leafy things?",
    "At 0.5 m the stripes are NOT canopy — they are cells where no first "
    "return came back",
    "3.23% of this window at 0.5 m.  0.00% at 1 m.",
    "Same laser, same ground. The half-metre grid was finer than the survey "
    "delivered.",
    "This kind of stripe is a grid-size problem, and one metre solves it. "
    "There is a second kind that it does not — slide 30.",
]
NOTE = """
Canopy height. How tall are the trees. You get it by taking the top surface and
subtracting the ground surface.

Now the important part, because people ask about it every time.

Look at the left panel. Those diagonal stripes are not trees and they are not a
rendering glitch. They are cells where no laser pulse came back at all, so there
is nothing to measure the height of.

Same window, same flight, three cell sizes. At half a metre, 3.23% of the cells
are empty. At one metre, zero.

Nothing was fixed or filled in. The half-metre grid was simply finer than the
survey actually delivered, so it asked for detail that was never measured.

That is why the terrain layers in this deck are built at one metre.

One thing to be precise about, because it comes up again shortly. These stripes
are empty cells. A grid-size problem, and one metre solves it completely.

There is a second kind of stripe, on the RRIM, that looks similar and is not the
same thing at all. Those sit in a surface with no empty cells anywhere, so
coarsening cannot delete them. That is slide 30.
"""

INK = RGBColor(0x14, 0x1A, 0x1F)
INK2 = RGBColor(0x54, 0x5C, 0x63)
ACCENT = RGBColor(0x00, 0x96, 0xC7)


def rgb_runs(tf, runs):
    first = True
    for text, size, bold, colour, space in runs:
        if first and len(tf.paragraphs) and not tf.paragraphs[0].runs:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Calibri"


def main() -> int:
    if not FIG.exists():
        raise SystemExit(f"figure missing: {FIG}")
    prs = Presentation(DECK)
    s = prs.slides[SLIDE - 1]

    txt = [sh for sh in s.shapes
           if sh.has_text_frame and sh.text_frame.text.strip()]
    if not txt or EXPECT not in txt[0].text_frame.text:
        raise SystemExit(f"slide {SLIDE} is not the canopy-height slide; "
                         f"found {txt[0].text_frame.text[:60]!r}")

    # drop the old 0.5 m picture and the old text block
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)

    body = s.shapes.add_textbox(Inches(0.55), Inches(0.30),
                                Inches(12.25), Inches(2.25))
    body.text_frame.word_wrap = True
    runs = [(TITLE, 26, True, ACCENT, 10)]
    for b in BULLETS:
        runs.append(("•  " + b, 15, False, INK2, 6))
    rgb_runs(body.text_frame, runs)

    # figure is 2550x1003, so 2.542:1. Full width, what is left of the height.
    w = 12.25
    s.shapes.add_picture(str(FIG), Inches(0.55), Inches(2.62),
                         width=Inches(w), height=Inches(w / 2.542))

    s.notes_slide.notes_text_frame.text = NOTE.strip()

    prs.save(DECK)
    print(f"  slide {SLIDE}: now shows 0.5 / 1 / 2 m side by side")
    print(f"  figure bottom at {2.62 + w / 2.542:.2f} in of 7.50")
    print(f"  {DECK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

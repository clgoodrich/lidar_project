"""Bring docProps/app.xml back in line with the slides that actually exist.

THE DEFECT
----------
`docProps/app.xml` carries a declared slide count, a "Slide Titles" heading
pair, and a `TitlesOfParts` vector listing every font, the theme, and one entry
per slide. PowerPoint writes it and checks it. **python-pptx never touches it.**

So every slide added or removed since the deck was last saved by PowerPoint
left it further out of step. In v9:

    app.xml says          <Slides>68</Slides>, TitlesOfParts vector size 71
    the deck actually has 75 slides
    PowerPoint's own repaired copy says 75, vector size 78

That mismatch is what the repair prompt is about. It survived every structural
check run against the package, because the package IS structurally valid --
the defect is a stale summary of it, in a part nothing else references.

It also explains why the prompt predates the recent work: the count has been
wrong since the first slide was added or deleted programmatically, which is
why v6, v7, v8 and v9 all prompt.

WHAT THIS WRITES
----------------
    <Slides>              the real slide count
    <Notes>               the real notes-slide count
    HeadingPairs          "Slide Titles" -> the real count
    TitlesOfParts         fonts + theme + one entry per slide, resized to match

Slide titles are taken from each slide's title placeholder where it has one,
and otherwise from its first non-empty text box -- which is how the rest of
this deck was built, so most slides have no real title placeholder. Where
neither exists the entry is "PowerPoint Presentation", matching what was
already there.

Run:
    python docs/presentation/_sync_docprops_app.py <in.pptx> [out.pptx]
Default output is the input with " SYNCED" before the extension.
"""
from __future__ import annotations

import shutil
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

EP = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
FALLBACK = "PowerPoint Presentation"

#: A soft line break inside a title comes back from python-pptx as \x0b. That
#: is legal in the slide (it is an <a:br/>) but not in app.xml, where lxml
#: refuses to write it at all. Slide 1's title carries one. PowerPoint writes a
#: space in its place, so this does the same.
import re as _re
_CTRL = _re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]+")


def clean(s: str) -> str:
    return _re.sub(r"\s+", " ", _CTRL.sub(" ", s)).strip()


def slide_title(slide) -> str:
    """The title PowerPoint would list for this slide."""
    # a real title placeholder first
    for ph in slide.placeholders:
        try:
            if ph.placeholder_format.type is not None and \
               "TITLE" in str(ph.placeholder_format.type):
                t = ph.text_frame.text.strip()
                if t:
                    return clean(t.split("\n")[0])
        except (AttributeError, ValueError):
            continue
    # otherwise the topmost non-empty text box, which is how this deck is built
    boxes = [sh for sh in slide.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    if boxes:
        top = min(boxes, key=lambda s: s.top if s.top is not None else 0)
        return clean(top.text_frame.text.strip().split("\n")[0])
    return FALLBACK


def main() -> int:
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    src = Path(sys.argv[1])
    dst = (Path(sys.argv[2]) if len(sys.argv) > 2
           else src.with_name(src.stem + " SYNCED" + src.suffix))

    prs = Presentation(src)
    titles = [slide_title(s) for s in prs.slides]
    n_slides = len(titles)
    n_notes = sum(1 for s in prs.slides if s.has_notes_slide)

    with zipfile.ZipFile(src) as z:
        app = etree.fromstring(z.read("docProps/app.xml"))

    def child(tag):
        for e in app:
            if etree.QName(e).localname == tag:
                return e
        return None

    before = (child("Slides").text if child("Slides") is not None else "?")

    # ---- scalar counts -------------------------------------------------
    for tag, val in (("Slides", n_slides), ("Notes", n_notes)):
        e = child(tag)
        if e is None:
            e = etree.SubElement(app, f"{{{EP}}}{tag}")
        e.text = str(val)

    # ---- how many fonts and themes the vector already declares ---------
    hp = child("HeadingPairs")
    vec = hp[0]
    counts = {}
    label = None
    for variant in vec:
        inner = variant[0]
        if etree.QName(inner).localname == "lpstr":
            label = inner.text
        else:
            counts[label] = int(inner.text)
    n_font = counts.get("Fonts Used", 0)
    n_theme = counts.get("Theme", 0)

    # "Slide Titles" is the only pair that moves
    label = None
    for variant in vec:
        inner = variant[0]
        if etree.QName(inner).localname == "lpstr":
            label = inner.text
        elif label == "Slide Titles":
            inner.text = str(n_slides)

    # ---- TitlesOfParts: keep the fonts and theme, replace the titles ----
    top = child("TitlesOfParts")
    tvec = top[0]
    keep = list(tvec)[:n_font + n_theme]
    for e in list(tvec):
        tvec.remove(e)
    for e in keep:
        tvec.append(e)
    for t in titles:
        el = etree.SubElement(tvec, f"{{{VT}}}lpstr")
        el.text = t
    tvec.set("size", str(n_font + n_theme + n_slides))

    new_app = etree.tostring(app, xml_declaration=True, encoding="UTF-8",
                             standalone=True)

    # ---- rewrite the package, swapping just this one part ---------------
    shutil.copy(src, dst)
    with zipfile.ZipFile(src) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == "docProps/app.xml":
                data = new_app
            zout.writestr(info, data)

    # ---- prove it -------------------------------------------------------
    with zipfile.ZipFile(dst) as z:
        chk = etree.fromstring(z.read("docProps/app.xml"))
    got_slides = [e.text for e in chk if etree.QName(e).localname == "Slides"][0]
    got_vec = [e[0].get("size") for e in chk
               if etree.QName(e).localname == "TitlesOfParts"][0]
    print(f"  slides in deck:        {n_slides}")
    print(f"  app.xml <Slides>:      {before}  ->  {got_slides}")
    print(f"  TitlesOfParts size:    {got_vec} "
          f"(= {n_font} fonts + {n_theme} theme + {n_slides} slides)")
    print(f"  notes slides:          {n_notes}")
    print(f"  {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

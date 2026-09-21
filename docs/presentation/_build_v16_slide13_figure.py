"""v15 -> v16: give slide 13 a chart that can be read.

WHAT WAS WRONG WITH THE OLD RIGHT-HAND PANEL
--------------------------------------------
Slide 13 carried a two-panel figure. The left panel is clear. The right one was
a dot plot that nobody should be expected to decode mid-talk:

  * the vertical position of every dot was pure jitter, carrying no meaning at
    all, while looking exactly like it encoded something
  * it used three marker shapes -- filled, hollow, and x -- whose meanings were
    split between a bold title and a grey note in the corner
  * the x markers sat at x = 0 to mean "there is NO cliff here", which reads as
    "the cliff is at zero degrees", the precise opposite
  * and its message -- that Venango has a gap between where ground stops and
    how wide they flew, and McKean does not -- is already the whole job of the
    NEXT slide, "one survey, not all of them"

So the panel was duplicating slide 14 in a form harder to read than slide 14.

WHAT REPLACES IT
----------------
The left panel alone, taken from the all-tiles version of the same chart, which
is a stronger figure than the six-square one that was embedded:

    Venango 2020-03   177 map squares   drops to 0% at 18.5 deg
    McKean  2019-04    59 map squares   tapers, never reaches zero
    Venango 2019-11    16 map squares   tapers to about 48%

252 map squares rather than 12, and three acquisitions rather than two. One
chart, one axis, one idea: of the returns that reached the ground, what share
the survey actually labelled as ground, against the angle they were shot at.

The figure's own footer note is cropped off and its content moved into the
slide's bullets, since text baked into a PNG cannot be re-wrapped or read at
presentation size.

Run:
    python docs/presentation/_build_v16_slide13_figure.py
Reads:  docs/presentation/WellSight_Presentation v15.pptx
Writes: docs/presentation/WellSight_Presentation v16.pptx
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs/presentation/WellSight_Presentation v15.pptx"
DST = ROOT / "docs/presentation/WellSight_Presentation v16.pptx"
FIG = (ROOT / "docs/presentation/figures_30to45min/1_data_qa"
       / "scan_angle_cliff_by_survey_all_tiles_leftonly.png")

SLIDE_TITLE = "Data QA — nothing past 18° was called ground"
INK = RGBColor(0x14, 0x1A, 0x1F)
INK2 = RGBColor(0x54, 0x5C, 0x63)

KICKER = "Of the returns that reached the ground, how many were labelled ground"
BULLETS = [
    "Each faint line is one map square; the thick line is the middle of its "
    "batch",
    "Venango March 2020 — 177 squares — falls to 0% at 18.5°. "
    "Not a taper, a wall.",
    "McKean April 2019 — 59 squares — tapers and never reaches zero",
    "Venango November 2019 — 16 squares — tapers to about 48%",
    "A clean vertical edge at a round number is a setting in software, never "
    "the physics",
    "Venango 2011 is excluded: those squares record no scan angle at all",
]
NOTES = (
    "One chart, one idea. Along the bottom is scan angle -- how far off "
    "straight-down the laser was pointing. Up the side is the share of returns "
    "that reached the ground and were actually labelled as ground.\n\n"
    "Every faint line is a single map square. The thick line is the middle of "
    "its batch, and the band holds the middle 80 percent.\n\n"
    "Blue is Venango, March 2020, 177 squares. It runs at 97 to 99 percent all "
    "the way out and then drops to zero at eighteen and a half degrees. "
    "Straight down. That is the whole point of the slide.\n\n"
    "Red is McKean, a year earlier, 59 squares. It sags to about 80 percent "
    "and recovers. No wall.\n\n"
    "Orange is Venango again but November 2019, only 16 squares. It tapers to "
    "about half. Also no wall.\n\n"
    "So this is not how airborne lidar behaves, and it is not how this county "
    "behaves. It is one batch of flights.\n\n"
    "If asked about the accuracy of what was cut: the discarded returns "
    "measure 6.5 to 6.7 cm against ground from neighbouring flight lines, "
    "versus 7.0 to 9.9 cm for the wide-angle returns that were kept. Both "
    "inside the 10 cm the specification allows.")


def rgb_runs(tf, runs):
    first = True
    for text, size, bold, colour, space in runs:
        if first and len(tf.paragraphs) and not tf.paragraphs[0].runs:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Calibri"


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def main() -> int:
    for p in (SRC, FIG):
        if not p.exists():
            raise SystemExit(f"missing: {p}")
    with Image.open(FIG) as im:
        ar = im.size[0] / im.size[1]

    prs = Presentation(SRC)
    slide = next((s for s in prs.slides
                  if title_of(s).strip() == SLIDE_TITLE), None)
    if slide is None:
        raise SystemExit(f"could not find {SLIDE_TITLE!r}")

    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)

    tb = slide.shapes.add_textbox(Inches(0.70), Inches(0.42),
                                  Inches(12.0), Inches(0.85))
    tb.text_frame.word_wrap = True
    rgb_runs(tb.text_frame, [(SLIDE_TITLE, 27, True,
                              RGBColor(0x00, 0x96, 0xC7), 0)])

    body = slide.shapes.add_textbox(Inches(0.70), Inches(1.50),
                                    Inches(4.55), Inches(5.10))
    body.text_frame.word_wrap = True
    runs = [(KICKER, 16, True, INK, 10)]
    for b in BULLETS:
        runs.append(("•  " + b, 13, False, INK2, 7))
    rgb_runs(body.text_frame, runs)

    max_w, max_h = 7.70, 5.35
    w = min(max_w, max_h * ar)
    h = w / ar
    slide.shapes.add_picture(str(FIG), Inches(5.45 + (max_w - w) / 2),
                             Inches(1.42 + (max_h - h) / 2),
                             width=Inches(w), height=Inches(h))

    foot = slide.shapes.add_textbox(Inches(0.70), Inches(7.00),
                                    Inches(11.0), Inches(0.42))
    foot.text_frame.word_wrap = True
    rgb_runs(foot.text_frame,
             [("Colton Goodrich  ·  University of Houston",
               12, False, INK2, 0)])
    slide.notes_slide.notes_text_frame.text = NOTES

    prs.save(DST)
    out = Presentation(DST)
    print(f"  slide rebuilt: {SLIDE_TITLE}")
    print(f"  figure {FIG.name} (ratio {ar:.2f}), one panel")
    print(f"  {len(out.slides)} slides")
    print(f"  {DST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

# -*- coding: utf-8 -*-
"""The one script that turns the user's current v15 into v18.

Re-run this after every edit the user makes to v15. It supersedes
_build_v17_reconcile_all_counts.py, _refresh_slide20_void_contrast_v17.py and
_fix_void_denominator_and_slide12_v17.py, which were three passes over the same
deck and are kept only as the record of how each change was reasoned.

Everything here is matched on slide TITLE or on exact paragraph text, never on
slide index, so the user inserting or deleting slides does not break it. Any
target that cannot be found stops the run rather than being skipped quietly.

--------------------------------------------------------------------------
1. THE COUNTS, EACH WITH ITS SCOPE
--------------------------------------------------------------------------
The Data QA section quoted counts from five scopes and named none of them,
which is why they read as contradictory:

    slide 10   113.5 million points          10 map squares
    slide 11   5.6 million                   4 map squares
    slide 12   190.9 million (in the figure) 165 of 177 squares, whole block
    slides 15-17, 21                         9t alone
    slides 18, 20                            one cross-section line, one window

Two were wrong, not merely unlabelled:

  * Slide 10's body says "nine tiles" while its own notes say "ten". The source
    is ten -- nonground_classification_and_scan_angle_cut.md line 30,
    113,556,364 points.
  * Slide 17 says "about a quarter of the holes close". It is 2,857,002 of
    7,098,612, which is 40.2%. Two in five.

--------------------------------------------------------------------------
2. WHY 13.8% IS CORRECT AND STAYS
--------------------------------------------------------------------------
An earlier pass replaced 13.8% with 45.2% and that was a mistake, corrected
here. Both numbers are real; they have different denominators.

    81,000,000   cells in the 9t bounding box at 0.5 m
    29,529,952   no return of ANY kind ever landed in them
    51,470,048   a return landed  <- the denominator that means something
    44,371,436   ...and it included a ground return
     7,098,612   ...and it did not          = 13.8%
     2,857,002   filled by restoring the cut = 40.2% of those, two in five
     4,241,610   left                        =  8.2%

44,371,436 + 7,098,612 = 51,470,048 exactly, so the summary JSON and the
rasters agree to the cell.

The 29.5 M never-sampled cells are a grid-spacing matter, not a vendor
decision: ground returns sit 0.61 m apart and the grid asks every 0.5 m.
Counting them as "no ground measurement" blames the survey for our own choice
of cell size. Definition source:
notebooks/wellsight_v2/s7_analysis/_recovered_ground_maps_9t.py line 767,
`inside = cnt_a > 0  # cells the laser actually reached`.

So 13.8% and 8.2% stand, and each now states its denominator on the slide --
which is the thing that was missing and what made 13.8% look irreconcilable
with every number around it.

--------------------------------------------------------------------------
3. THE THREE FIGURES
--------------------------------------------------------------------------
slide 12  the card stack, rebuilt so the 191 million card carries its
          per-square rate (1.16 M across 165 squares). Without it nobody can
          see that 191 M and slide 10's 113.5 M describe the same survey.
slide 14  the flight-block bars without the Venango 2011-09 row. That delivery
          never populated the scan angle, so its bar meant "cannot be checked"
          among bars that mean a measured angle.
slide 20  the void class now separates by hue AND by contrast. A flat orange
          wash at alpha 0.30 did not survive 52% coverage: the whole panel went
          warm and the boundary vanished.

--------------------------------------------------------------------------
4. SLIDE 57
--------------------------------------------------------------------------
Speaker notes gain the 128 m training patch and the 30 m jitter, appended to
whatever note is already there.

Run:
    python docs/presentation/_build_v18_from_v15.py
Reads:  docs/presentation/WellSight_Presentation v15.pptx
Writes: docs/presentation/WellSight_Presentation v18.pptx
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
from _annotation_slides_9t_only import (  # noqa: E402
    ANNOTATION_EDITS, ANNOTATION_STALE,
)
from _speaker_note_additions import NOTE_APPENDS  # noqa: E402
from _separate_the_two_evaluations import (  # noqa: E402
    EVAL_EDITS, EVAL_STALE,
)
from _port_v17_user_edits import (  # noqa: E402
    V17_DELETED_SLIDES, V17_EDITS, V17_NOTE_APPENDS, V17_NOTE_REPLACE,
    V17_RETIRED_NOTES,
)
from _side_text_613590_roads import ROAD_SIDE_TEXT  # noqa: E402
from _number_audit_fixes import AUDIT_EDITS, AUDIT_STALE  # noqa: E402
from _american_spellings import audit as spell_audit  # noqa: E402
from _american_spellings import normalise_deck  # noqa: E402
from _nisar_broad_statements import (  # noqa: E402
    NISAR_ANCHOR, NISAR_EDITS, NISAR_NEW_BULLETS, NISAR_NOTE, NISAR_TITLE,
)

HERE = Path(__file__).resolve().parent
SRC = HERE / "WellSight_Presentation v15.pptx"
DST = HERE / "WellSight_Presentation v18.pptx"
FIGDIR = HERE / "figures_30to45min"
PLANVIEW = FIGDIR / "v6/median_pit_plan_view_9t_05.png"
SPLITKEY = FIGDIR / "v6/block_split_key_9t.png"
PITDETAIL = FIGDIR / "v6/pit_probability_detail_880x480m_9t_05.png"
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

PITDETAIL_AFTER = "Outcome – Pits"
PITDETAIL_TITLE = "Outcome – Pits, up close"
PITDETAIL_NOTE = (
    "The same thing as the slide before, given a whole slide so it can "
    "actually be read.\n\n"
    "880 by 480 metres of ground. Grey underneath is the shaded terrain. The "
    "bright colour is the model's probability that a pixel is a pit floor "
    "— dark purple is low, orange and yellow are high. Anything under "
    "0.05 is left off entirely so the terrain shows through.\n\n"
    "Blue outlines are the pit floors we drew by hand. They are the answer "
    "key, not a model output.\n\n"
    "What to look at: nearly every blue outline has a bright core sitting "
    "inside it, and the bright cores are small and sharp rather than smeared. "
    "The model is not painting the whole hillside amber and hoping.\n\n"
    "And the ground here was never used for training. Every pixel on this "
    "slide was predicted by whichever of the five models had this block held "
    "out, so none of them had seen this terrain when they made the "
    "prediction. 30 annotated floors are in view.\n\n"
    "If asked about the faint marks away from the outlines: those are the "
    "model responding weakly to other hollows. That is what the probability "
    "cut-off is for, and it is why we report a threshold rather than a "
    "picture."
)

#: The eight RRIM-beside-probability figures lost the two grey caption lines
#: that used to sit above the panels. Baked-in caption text cannot be
#: re-wrapped, cannot be read from the back of a room, and half of it repeated
#: the slide title. The captions were load-bearing though -- the crop IS a best
#: case by construction -- so the wording moves into the speaker notes here.
#: title -> (figure, task, operating threshold, tile)
PROB_FIGS = {
    "Outcome - Roads":              ("road", 0.20, "9t"),
    "Outcome - Drainage":           ("drainage", 0.50, "9t"),
    "Outcome - Pads":               ("pad", 0.45, "9t"),
    "Outcome – Pits":          ("pit", 0.20, "9t"),
    "Second tile, 613590 – Pits":     ("pit", 0.20, "613590"),
    "Second tile, 613590 – Pads":     ("pad", 0.45, "613590"),
    "Second tile, 613590 – Roads":    ("road", 0.20, "613590"),
    "Second tile, 613590 – Drainage": ("drainage", 0.50, "613590"),
}
PROB_DIR = FIGDIR / "5_probability_surfaces/v6"


def prob_caveat(task, thr, tile):
    seen = ("This is the tile the models trained on."
            if tile == "9t" else
            "The models never saw this tile. Nothing here was trained on.")
    return (
        "Two things to say about this figure, because they used to be printed "
        "on it and are not any more.\n\n"
        f"{seen}\n\n"
        f"The window was not chosen because it looks good. For every task the "
        f"script bins the pixels above the operating threshold — "
        f"{thr:.2f} for {task} — and takes the densest 400 m box, with the "
        "same rule applied to both tiles. It is still a best case. It is a "
        "best case found by a rule rather than by taste, and that is the "
        "honest way to show one.\n\n"
        "The right panel is drawn the way QGIS draws it: black is 0, white is "
        "1, nothing masked. So the near-black background is real low "
        "probability, not missing data."
    )


#: The outcome slides keep their probability maps untouched and gain a small
#: key in the left margin showing which blocks were train, validation and test.
#: Drawn on the same extent, so a reader can look straight across.
KEY_SLIDES = ["Outcome - Roads", "Outcome - Drainage", "Outcome - Pads",
              "Outcome – Pits"]
KEY_NOTE = (
    "The small key on the left says which blocks were which. Blue is train, "
    "the model learned from those. Orange hatched is validation, used to "
    "decide when to stop and where to set the cut-off. Red crossed is test, "
    "which nothing touched until the end. It is drawn on the same extent as "
    "the map, so a square here is the same square there."
)

BULLET = "•  "
EMDASH = "—"

#: (old, new, "body" | "notes"). Matched exactly; a miss aborts the run.
EDITS = [
    # -- slide 10: body and notes disagreed on the count -------------------
    (BULLET + "113.5 million points across nine tiles",
     BULLET + "113.5 million points across ten map squares",
     "body"),
    ("The delivery is 113.5 million points over ten tiles.",
     "The delivery is 113.5 million points over ten map squares. That is the "
     "audit set, wider than the nine squares of 9t that the rest of the talk "
     "uses.",
     "notes"),
    # -- slide 11: name the scope ------------------------------------------
    (BULLET + "5.6 million of them, across four tiles",
     BULLET + "5.6 million of them, across four of those map squares",
     "body"),
    ("There are 5.6 million of them across four tiles.",
     "There are 5.6 million of them across four of those map squares.",
     "notes"),
    # -- slide 15: keep 13.8%, say what it is a share OF --------------------
    (BULLET + "13.8% of the training area has nothing under it.",
     BULLET + "13.8% of the ground the laser reached has no measurement on it",
     "body"),
    # -- slide 16: correct, but say where ----------------------------------
    (BULLET + "13.7 million ground measurements.",
     BULLET + "13.7 million ground measurements, inside 9t alone",
     "body"),
    ("13.7 million ground measurements.",
     "13.7 million ground measurements, inside 9t alone.",
     "notes"),
    # -- slide 17: keep the rates, fix the fraction ------------------------
    (BULLET + "13.8% with no measurement, down to 8.2%.",
     BULLET + "13.8% with no ground measurement, down to 8.2%",
     "body"),
    (BULLET + "2.9 million cells filled in.",
     BULLET + "2.9 million of those 7.1 million holes filled in",
     "body"),
    (BULLET + "About a quarter of the holes close. The rest are canopy.",
     BULLET + "Two holes in five close. The rest is canopy the laser never "
     "got through.",
     "body"),
    # -- slide 21: correct, but say the grid -------------------------------
    (BULLET + "44.4 million cells already had ground. 2.9 million were filled.",
     BULLET + "In 9t at 0.5 m: 44.4 million cells already had ground, 2.9 "
     "million were filled",
     "body"),
    # -- slides 34-36: the corn rows are NOT tens of centimetres -----------
    # Measured twice, independently: 0.91 cm across the nine hand-drawn rows
    # (_cornrow_rows_measured_9t.py) and 0.70 cm on the cross-strike profile
    # (_cornrow_lrm_cross_strike_9t.py), both 5-95%. Slide 37 already said
    # "under a centimetre" and was right. "Tens of centimetres" was out by
    # roughly thirty times.
    (BULLET + "Their size is tens of centimetres. A pit is a 0.7 m dish.",
     BULLET + "They stand under a centimetre tall. A pit is a 0.7 m dish.",
     "body"),
    # The old claim followed from the wrong number: a 0.9 cm ripple and a
    # 0.7 m pit are not the same size, they are seventy times apart. What IS
    # true is that they occupy the same SPACING, and that the layers the model
    # reads respond to shape rather than to height.
    (BULLET + "So the artefact and the signal we are hunting are the same "
     "size. That is why it matters.",
     BULLET + "They are far shallower than a pit, but they repeat every "
     "3.35 m — the width of the pit rims we are looking for",
     "body"),
    ("Their size is tens of centimetres. A collapse pit is a dish about "
     "0.7 m deep.",
     "They stand under a centimetre tall, measured two ways: 0.91 cm across "
     "the nine rows drawn by hand, 0.70 cm on the cross-strike profile. A "
     "collapse pit is a dish about 0.7 m deep, so the rows are far shallower.",
     "notes"),
    ("So the artefact and the target are the same size, which is why it "
     "matters here",
     "They are not the same depth as a pit. What they share is spacing: a "
     "crest every 3.35 m, which is the width of the rims we look for. And "
     "every layer the model reads -- local relief, openness, hillshade -- "
     "measures shape, not height, so a ripple this shallow still prints on "
     "them",
     "notes"),
    ("that layer is elevation with the hillside subtracted, so a few "
     "centimetres of ripple is most of what is left.",
     "that layer is elevation with the hillside subtracted, so a ripple of "
     "well under a centimetre is most of what is left.",
     "notes"),
]

EDITS += ANNOTATION_EDITS
EDITS += NISAR_EDITS
EDITS += EVAL_EDITS
EDITS += V17_EDITS
# Last, because several of these correct the OUTPUT of an edit above
# (the 0.7 m dish bullet, the corn-row note) rather than the v15 text.
EDITS += AUDIT_EDITS

#: title -> figure. Position and box are inherited from the picture replaced.
#: "fit" reflows the height to the new aspect instead of stretching.
FIGURES = {
    "Data QA – What the Cut Cost":
        (FIGDIR / "1_data_qa/data_qa_scan_angle_cut_numbers_9t.png", "keep"),
    "Data QA — one survey, not all of them":
        (FIGDIR / "v6/where_ground_stops_by_flight_block_excl2011_9t.png",
         "fit"),
    "Data QA — the surface, built both ways":
        (FIGDIR / "v6/dem_with_and_without_deleted_returns_9t.png", "keep"),
}

SLIDE_57_TITLE = "Model building — does a bigger network help?"
PATCH_NOTE = (
    "If asked what a training patch is:\n\n"
    "The tile is too big to feed a network at once, so training happens on "
    "small square cut-outs. Each one is 256 pixels at 0.5 m, which is 128 m on "
    "the ground.\n\n"
    "Patches are centred on annotated pits, otherwise almost every patch would "
    "be empty forest.\n\n"
    "The jitter is the important half. Without it the pit sits dead centre in "
    "every single patch, and the network can learn “the answer is in the "
    "middle” instead of what a pit looks like. So each patch centre is "
    "shifted by a random amount up to 30 m before it is cut. The pit lands "
    "somewhere different every time it is seen.\n\n"
    "One guard goes with it. A jittered patch is clipped to the training "
    "blocks, so a shifted patch cannot reach across the line into the held-out "
    "fold."
)

#: text that must not survive, from either the original deck or a bad pass
STALE = ["across nine tiles", "across four tiles", "a quarter of the holes",
         "45.2%", "41.7%", "one hole in thirteen", "tens of centimetres",
         "are the same size"]
STALE += ANNOTATION_STALE
STALE += EVAL_STALE
STALE += AUDIT_STALE


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def edit_body(slide, old, new):
    """Replace a paragraph's text, or DELETE the paragraph when new is None.

    Deleting matters for the annotation slides: the per-region split lines
    ("9t ... | 613590 ... | elsewhere ...") exist to break a total into
    regions, and once the slide is one region there is no total to break.
    """
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            if p.text.strip() != old.strip():
                continue
            if new is None:
                p._p.getparent().remove(p._p)
                return True
            runs = list(p.runs)
            if not runs:
                continue
            runs[0].text = new
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            return True
    return False


def edit_notes(slide, old, new):
    if not slide.has_notes_slide:
        return False
    tf = slide.notes_slide.notes_text_frame
    if old not in tf.text:
        return False
    tf.text = tf.text.replace(old, new, 1)
    return True


#: Ten more figures rebuilt with WELLSIGHT_BARE=1, which strips the headline,
#: subtitle and attribution baked into the PNG. The slide already carries a
#: title; a second one inside the image cannot be re-wrapped, cannot be read
#: from the back of a room, and on the road-network figure the right-hand note
#: printed straight through the title beside it.
#: title -> figure, all placed with mode "contain" because removing the title
#: band changes the aspect and the existing box must not stretch them.
BARE_FIGS = {
    "Manual Annotation — Roads":
        "3_annotations/v6/annotation_roads_2km_lines_9t.png",
    "Manual Annotation — Drainage":
        "3_annotations/v6/annotation_drainage_2km_lines_9t.png",
    "Manual Annotation — Pads":
        "3_annotations/v6/annotation_pads_2km_lines_9t.png",
    "Manual Annotation – Pits (Outside)":
        "3_annotations/v6/annotation_pits_1km_rims_9t.png",
    "Manual Annotation – Pits (Inside)":
        "3_annotations/v6/annotation_pits_1km_floors_9t.png",
    "Preprocessing — matching rims to floors":
        "3_annotations/v6/pit_rim_floor_matching_9t.png",
    "Second tile, 613590 — stuff to investigate":
        "5_probability_surfaces/v6/pit_pad_candidates_613590.png",
    "Second tile, 613590 – The Road Network It Drew":
        "5_probability_surfaces/v6/roads_generated_thr0p30_613590.png",
    "Second tile, 613590 – Generated Roads vs TIGER":
        "5_probability_surfaces/v6/roads_vs_tiger_613590.png",
    "Second tile, 613590 – Roads Found and Missed":
        "5_probability_surfaces/v6/roads_found_vs_missed_thr0p50_613590.png",
}


def swap_figure(slide, png, mode):
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if len(pics) != 1:
        raise SystemExit(f"{title_of(slide)!r} has {len(pics)} pictures, "
                         "expected exactly 1")
    old = pics[0]
    left, top, width, height = old.left, old.top, old.width, old.height
    with Image.open(png) as im:
        w, h = im.size
    if mode == "fit":
        height = Emu(int(width / (w / h)))
    elif mode == "contain":
        # Largest size that fits inside the old box, centred in it. Stripping
        # a title band changes the aspect, and reusing the box unchanged would
        # stretch the map.
        ar = w / h
        if width / height > ar:
            new_w = Emu(int(height * ar))
            left = Emu(int(left + (width - new_w) / 2))
            width = new_w
        else:
            new_h = Emu(int(width / ar))
            top = Emu(int(top + (height - new_h) / 2))
            height = new_h
    elif abs((w / h) - (width / height)) / (w / h) > 0.02:
        raise SystemExit(f"{png.name}: aspect changed and mode is 'keep'; "
                         "switch to 'fit' or reposition deliberately")
    old_sha = old.image.sha1[:12]
    old._element.getparent().remove(old._element)
    new = slide.shapes.add_picture(str(png), left, top,
                                   width=width, height=height)
    return old_sha, new.image.sha1[:12], width, height


def main() -> int:
    for p in [SRC] + [f for f, _ in FIGURES.values()]:
        if not p.exists():
            raise SystemExit(f"missing: {p}")

    prs = Presentation(SRC)
    n_before = len(prs.slides)
    slides = list(prs.slides)
    print(f"  reading v15: {n_before} slides\n")

    # ---- 1. the counts ---------------------------------------------------
    misses = []
    for old, new, where in EDITS:
        hit = 0
        for s in slides:
            if where == "body" and edit_body(s, old, new):
                hit += 1
            elif where == "notes" and edit_notes(s, old, new):
                hit += 1
        if hit:
            shown = "(paragraph deleted) " + old[:52] if new is None else new[:72]
            print(f"  ok  [{where:5s}] {shown}")
        else:
            misses.append(f"[{where}] {old[:68]}")
    if misses:
        raise SystemExit("targets not found in v15, so it has moved:\n  "
                         + "\n  ".join(misses))

    # ---- 2. the figures --------------------------------------------------
    print()
    for tl, (png, mode) in FIGURES.items():
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        a, b, wd, ht = swap_figure(hits[0], png, mode)
        print(f"  fig {a} -> {b}  {png.name}")
        print(f"      {wd/914400:.2f} x {ht/914400:.2f} in  ({mode})")

    # ---- 2b. slide 46 gains the plan view, stacked above the section ----
    # Slide 46 had a wide, short cross-section banner and 5.35 in of clear
    # space between the subtitle and the footer. The plan view goes on top at
    # full readable width and the section banner is shrunk to sit under it,
    # rather than adding a slide and splitting the morphology across two.
    s46 = [s for s in slides if title_of(s).strip() == "Measured pit morphology"]
    if len(s46) != 1:
        raise SystemExit(f"expected one morphology slide, found {len(s46)}")
    s46 = s46[0]
    pics = [sh for sh in s46.shapes if sh.shape_type == 13]
    if len(pics) == 1:
        sec = pics[0]
        sec.left, sec.top = Emu(int(2.65 * 914400)), Emu(int(5.62 * 914400))
        sec.width, sec.height = Emu(int(8.00 * 914400)), Emu(int(1.49 * 914400))
        with Image.open(PLANVIEW) as im:
            ar = im.size[0] / im.size[1]
        w = 9.20
        s46.shapes.add_picture(str(PLANVIEW),
                               Emu(int((13.33 - w) / 2 * 914400)),
                               Emu(int(1.98 * 914400)),
                               width=Emu(int(w * 914400)),
                               height=Emu(int(w / ar * 914400)))
        print(f"\n  slide 46 gains {PLANVIEW.name} at {w:.2f} x "
              f"{w/ar:.2f} in; section banner moved below it")
    else:
        print(f"\n  slide 46 already has {len(pics)} pictures, left alone")

    # ---- 2a2. the de-captioned probability figures ----------------------
    print()
    for tl, (task, thr, tile) in PROB_FIGS.items():
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        png = PROB_DIR / f"rrim_vs_prob_{task}_400m_{tile}.png"
        if not png.exists():
            raise SystemExit(f"missing: {png}")
        a, b, wd, ht = swap_figure(hits[0], png, "keep")
        tf = hits[0].notes_slide.notes_text_frame
        note = prob_caveat(task, thr, tile)
        if "used to be printed" not in tf.text:
            tf.text = tf.text.rstrip() + "\n\n" + note
        print(f"  {tl[:34]:36s} {a} -> {b}  {png.name}")

    # ---- 2a3. the ten de-titled figures ---------------------------------
    print()
    for tl, rel in BARE_FIGS.items():
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        png = FIGDIR / rel
        if not png.exists():
            raise SystemExit(f"missing: {png}  (run the builder with "
                             "WELLSIGHT_BARE=1)")
        a, b, wd, ht = swap_figure(hits[0], png, "contain")
        print(f"  {tl[:40]:42s} {a} -> {b}  "
              f"{wd/914400:.2f}x{ht/914400:.2f} in")

    # ---- 2a4. side text for the three 613590 road slides ----------------
    # The captions came off the top of those PNGs; they go back as slide text
    # in the left column, which is where they belong.
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    print()
    for tl, (kicker, bullets) in ROAD_SIDE_TEXT.items():
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        box = next((sh for sh in hits[0].shapes
                    if sh.has_text_frame
                    and sh.text_frame.text.strip().startswith(tl)), None)
        if box is None:
            raise SystemExit(f"no title box on {tl!r}")
        tf = box.text_frame
        if kicker in tf.text:
            print(f"  side text already there: {tl[:44]}")
            continue
        tf.word_wrap = True
        for text, size, bold, colour, space in (
                [("", 8, False, 0x54_5C_63, 0),
                 (kicker, 16, True, 0x14_1A_1F, 10)]
                + [(BULLET + b, 13, False, 0x54_5C_63, 8) for b in bullets]):
            para = tf.add_paragraph()
            para.space_after = Pt(space)
            run = para.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(colour >> 16, (colour >> 8) & 255,
                                          colour & 255)
            run.font.name = "Calibri"
        print(f"  side text +{len(bullets)} bullets: {tl[:44]}")

    # ---- 2b2. the block-split key on the outcome slides -----------------
    print()
    for tl in KEY_SLIDES:
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        sl = hits[0]
        pics = [sh for sh in sl.shapes if sh.shape_type == 13]
        if len(pics) != 1:
            print(f"  {tl}: {len(pics)} pictures already, key not added")
            continue
        with Image.open(SPLITKEY) as im:
            ar = im.size[0] / im.size[1]
        kw = 1.58
        sl.shapes.add_picture(str(SPLITKEY), Emu(int(0.07 * 914400)),
                              Emu(int(1.78 * 914400)),
                              width=Emu(int(kw * 914400)),
                              height=Emu(int(kw / ar * 914400)))
        tf = sl.notes_slide.notes_text_frame
        if "small key on the left" not in tf.text:
            tf.text = tf.text.rstrip() + "\n\n" + KEY_NOTE
        print(f"  {tl}: key added at {kw:.2f} x {kw/ar:.2f} in")

    # ---- 2b3. a full-slide pit probability map, inserted after Outcome Pits
    # The outcome slide gives probability half a slide beside the RRIM, which
    # is too small to read anything but "some blobs exist". This adds a second
    # slide that is nothing but the surface, at 880 x 480 m, every pixel
    # out-of-fold. Cloned from the slide it follows so it inherits the theme;
    # only text shapes are copied, because a deep-copied picture element
    # points at a relationship the new slide does not own.
    src = [s for s in slides if title_of(s).strip() == PITDETAIL_AFTER]
    if len(src) != 1:
        raise SystemExit(f"expected one {PITDETAIL_AFTER!r}, found {len(src)}")
    src = src[0]
    had_detail = any(title_of(s).strip() == PITDETAIL_TITLE for s in slides)
    if had_detail:
        print(f"\n  {PITDETAIL_TITLE!r} already present, not re-added")
    else:
        import copy as _copy
        new = prs.slides.add_slide(src.slide_layout)
        for shp in src.shapes:
            if shp.shape_type == 13:
                continue
            new.shapes._spTree.append(_copy.deepcopy(shp._element))
        for shp in list(new.shapes):
            if not shp.has_text_frame:
                continue
            t = shp.text_frame.text.strip()
            if t.split("\n")[0] == title_of(src).strip():
                for p in shp.text_frame.paragraphs:
                    if p.runs:
                        p.runs[0].text = PITDETAIL_TITLE
                        for r in p.runs[1:]:
                            r._r.getparent().remove(r._r)
                        break
        with Image.open(PITDETAIL) as im:
            ar = im.size[0] / im.size[1]
        # Fit by HEIGHT. The slide is 7.5 in tall and the title takes the
        # top, so a full-width placement would run 0.35 in off the bottom.
        ph, top = 5.95, 1.22
        pw = min(12.45, ph * ar)
        ph = pw / ar
        new.shapes.add_picture(str(PITDETAIL),
                               Emu(int((13.33 - pw) / 2 * 914400)),
                               Emu(int(top * 914400)),
                               width=Emu(int(pw * 914400)),
                               height=Emu(int(ph * 914400)))
        new.notes_slide.notes_text_frame.text = PITDETAIL_NOTE
        # move it from the end into position directly after its source
        lst = prs.slides._sldIdLst
        ids = list(lst)
        at = next(i for i, sid in enumerate(ids)
                  if prs.part.rels[sid.get(RID)].target_part is src.part)
        lst.remove(ids[-1])
        lst.insert(at + 1, ids[-1])
        print(f"\n  inserted {PITDETAIL_TITLE!r} after {PITDETAIL_AFTER!r}")
        print(f"     {PITDETAIL.name} at {pw:.2f} x {ph:.2f} in")
        slides = list(prs.slides)

    # ---- 2b4. the NISAR slide gains two capability bullets ---------------
    # Everything already on that slide is about whether the instrument works.
    # Nothing said what it would be FOR, which is the half a reviewer funds.
    ns = [s for s in slides if title_of(s).strip() == NISAR_TITLE]
    if len(ns) != 1:
        raise SystemExit(f"expected one {NISAR_TITLE!r}, found {len(ns)}")
    ns = ns[0]
    body = None
    last = None
    for sh in ns.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            if p.text.strip() == NISAR_ANCHOR.strip() and p.runs:
                body, last = sh, p
    if body is None:
        raise SystemExit(f"anchor bullet not found on the NISAR slide:\n  "
                         f"{NISAR_ANCHOR}")
    if NISAR_NEW_BULLETS[0][:30] in body.text_frame.text:
        print(f"\n  {NISAR_TITLE}: capability bullets already present")
    else:
        import copy as _copy
        for text in NISAR_NEW_BULLETS:
            new_p = _copy.deepcopy(last._p)          # inherits the styling
            last._p.addnext(new_p)
            from pptx.text.text import _Paragraph
            para = _Paragraph(new_p, last._parent)
            para.runs[0].text = "•  " + text
            for r in para.runs[1:]:
                r._r.getparent().remove(r._r)
            last = para
        print(f"\n  {NISAR_TITLE}: +{len(NISAR_NEW_BULLETS)} capability "
              "bullets")
    ns.notes_slide.notes_text_frame.text = NISAR_NOTE

    # ---- 2c. speaker-note additions -------------------------------------
    print()
    for tl, marker, text in NOTE_APPENDS:
        if tl in V17_RETIRED_NOTES:
            print(f"  skipped, slide deleted in v17: {tl}")
            continue
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        tf = hits[0].notes_slide.notes_text_frame
        if marker in tf.text:
            print(f"  note already present: {tl}")
            continue
        tf.text = tf.text.rstrip() + "\n\n" + text
        print(f"  note +{len(text):4d} chars: {tl}")

    # ---- 2d. the user's own v17 edits ------------------------------------
    print()
    for tl, marker, text in V17_NOTE_APPENDS:
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        tf = hits[0].notes_slide.notes_text_frame
        if marker in tf.text:
            print(f"  v17 note already there: {tl[:44]}")
            continue
        tf.text = tf.text.rstrip() + "\n\n" + text
        print(f"  v17 note +{len(text):4d}: {tl[:44]}")

    for tl, text in V17_NOTE_REPLACE.items():
        hits = [s for s in slides if title_of(s).strip() == tl]
        if len(hits) != 1:
            raise SystemExit(f"expected one {tl!r}, found {len(hits)}")
        hits[0].notes_slide.notes_text_frame.text = text
        print(f"  v17 note REPLACED: {tl}")

    # Deletions go last, so nothing above has to care that a slide is gone.
    lst = prs.slides._sldIdLst
    for tl in V17_DELETED_SLIDES:
        doomed = [s for s in slides if title_of(s).strip() == tl]
        if len(doomed) != 1:
            raise SystemExit(f"expected one {tl!r} to delete, "
                             f"found {len(doomed)}")
        for sid in list(lst):
            if prs.part.rels[sid.get(RID)].target_part is doomed[0].part:
                lst.remove(sid)
                prs.part.drop_rel(sid.get(RID))
                break
        print(f"  v17 slide DELETED: {tl}")
    n_deleted = len(V17_DELETED_SLIDES)
    slides = list(prs.slides)

    # ---- 3. slide 57 -----------------------------------------------------
    s57 = [s for s in slides if title_of(s).strip() == SLIDE_57_TITLE]
    if len(s57) != 1:
        raise SystemExit(f"expected one {SLIDE_57_TITLE!r}, found {len(s57)}")
    tf = s57[0].notes_slide.notes_text_frame
    if "jitter" in tf.text.lower():
        print("\n  slide 57 already explains jitter, left alone")
    else:
        tf.text = tf.text.rstrip() + "\n\n" + PATCH_NOTE
        print("\n  slide 57 notes gain the 128 m patch and 30 m jitter")

    # ---- 4. spellings ----------------------------------------------------
    # Last, so it catches the text every step above wrote as well as v15's own.
    # v15 keeps its spellings; it is the user's file and is never written to.
    n_spell = normalise_deck(prs)
    left = sorted(set(spell_audit(prs)))
    print(f"\n  {n_spell} runs moved to American spellings")
    if left:
        print("  STILL BRITISH (split across runs, fix by hand):")
        for i, b in left:
            print(f"    slide {i}: {b}")

    prs.save(DST)

    # ---- verify ----------------------------------------------------------
    out = list(Presentation(DST).slides)
    expect = n_before + (0 if had_detail else 1) - n_deleted
    if len(out) != expect:
        raise SystemExit(f"slide count {n_before} -> {len(out)}, expected "
                         f"{expect}; not intended")
    stale = []
    for i, s in enumerate(out, 1):
        txt = "\n".join(sh.text_frame.text for sh in s.shapes
                        if sh.has_text_frame)
        stale += [f"slide {i}: {bad}" for bad in STALE if bad in txt]
    print(f"\n  {len(out)} slides in, {len(out)} out")
    print(f"  stale text left on a slide: {stale if stale else 'none'}")
    print(f"  {DST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

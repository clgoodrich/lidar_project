"""Put the hand-annotation counts onto the five Manual Annotation slides.

WHAT IT DOES
------------
v6 slides 30, 31, 32, 34 and 35 carry a full-bleed figure on the right and a
left column holding nothing but a title. This fills that column: how much was
drawn, in how many pieces, and how big the average one is. The same numbers go
into the speaker notes in plain language.

WHERE THE NUMBERS COME FROM
---------------------------
docs/presentation/figures_30to45min/annotation_inventory_9t_613590.json, written
by _annotation_inventory_stats.py in that folder. Nothing here is typed by hand,
so re-running the stats script and then this one keeps the slides honest.

TOTALS ARE NOT 9t TOTALS. Roads, pads and pits all run well past the 9t tile --
roads reach McKean. Quoting one grand total on a slide about the training tile
is the kind of number that draws a question from the floor with no good answer,
so each slide gives the total and then splits it 9t / 613590 / elsewhere.

Run:
    python docs/presentation/_add_annotation_inventory_to_v6.py
Edits in place:
    docs/presentation/WellSight_Presentation v6.pptx
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
DECK = HERE / "WellSight_Presentation v6.pptx"
STATS = HERE / "figures_30to45min" / "annotation_inventory_9t_613590.json"

INK = RGBColor(0x14, 0x1A, 0x1F)
INK2 = RGBColor(0x54, 0x5C, 0x63)

#: The five slides this touches, by v6 position.
SLIDES = (30, 31, 32, 34, 35)


def f(x, n=0):
    return f"{x:,.{n}f}"


def dia(area_m2):
    """Diameter of a circle of this area -- the only shape an audience pictures
    from a square-metre figure without doing arithmetic in their head."""
    return (4.0 * area_m2 / 3.14159265) ** 0.5


def build(stats):
    L = {d["layer"]: d for d in stats["lines"]}
    P = {d["layer"]: d for d in stats["polygons"]}

    def area_line(d, key):
        b = d["by_area"]
        u = "km" if key == "length_km" else "ha"
        return (f"9t {f(b['9t'][key], 1)} {u}   |   "
                f"613590 {f(b['613590'][key], 1)} {u}   |   "
                f"elsewhere {f(b['elsewhere'][key], 1)} {u}")

    r, nr, dr = L["roads"], L["not_roads"], L["drainage"]
    pads = P["plat"]
    out, ins, wall = P["pit_outside"], P["pit_inside"], P["pit_wall"]
    spec = {}

    spec[30] = dict(
        kicker=f"{f(r['total_length_km'], 1)} km of road, drawn by hand",
        bullets=[
            f"{f(r['n_segments'])} separate lines, joining into "
            f"{f(r['n_networks'])} connected networks",
            f"Median line {f(r['median_segment_m'])} m, "
            f"mean {f(r['mean_segment_m'])} m",
            area_line(r, "length_km"),
            f"Plus {f(nr['n_segments'])} lines, {f(nr['total_length_km'], 1)} km, "
            f"drawn as not-road",
        ],
        notes=f"""HOW MUCH ROAD WENT IN BY HAND

{f(r['total_length_km'], 1)} kilometres of it. Somebody sat in QGIS and traced
every one of those lines over the terrain images.

{f(r['n_segments'])} separate strokes. They join end to end into
{f(r['n_networks'])} connected networks, which is closer to how many actual
roads there are, because one road usually takes several strokes to draw.

A typical stroke is {f(r['median_segment_m'])} m long. The average is
{f(r['mean_segment_m'])} m, pulled up by a few long hauls.

WHERE IT IS

{area_line(r, 'length_km')}

The elsewhere is Oil Creek through to McKean. It is drawn, but no model has
been trained on it yet.

THE NOT-ROAD LAYER

{f(nr['n_segments'])} more lines, {f(nr['total_length_km'], 1)} km. These are
things that look like a road in the imagery and are not one: old grades,
drainage cuts, field edges. They go to the model as negatives, so it learns the
difference during training instead of being corrected afterwards by a filter."""
    )

    cls = dr.get("by_class_km", {})
    spec[31] = dict(
        kicker=f"{f(dr['total_length_km'], 1)} km of drainage, 9t only",
        bullets=[
            f"{f(dr['n_drawn_lines'])} drawn lines, cut into "
            f"{f(dr['n_segments'])} training segments",
            f"Median segment {f(dr['median_segment_m'])} m, "
            f"mean {f(dr['mean_segment_m'])} m",
            "   |   ".join(f"{k} {v:.1f} km" for k, v in sorted(cls.items())),
            f"{f(dr['n_networks'])} connected networks",
        ],
        notes=f"""HOW MUCH DRAINAGE WENT IN BY HAND

{f(dr['total_length_km'], 1)} kilometres, all of it on 9t. Nothing was drawn on
613590, which is why the drainage class only ever trains on one tile.

{f(dr['n_drawn_lines'])} lines were drawn. Those got chopped into
{f(dr['n_segments'])} short segments for training, about
{f(dr['median_segment_m'])} m each. The chopping is deliberate. A model that
sees a whole 400 m stream at once learns "long thin thing", and then calls every
long thin thing a stream.

WHY DRAINAGE IS ANNOTATED AT ALL

It is not a product anybody asked for. A stream cut and an access road look
almost identical on a slope image. Both are long, both are darker than the
ground either side. Drawing the streams and handing them to the road model as
negatives is what stopped it calling them roads. That was fixed in training,
not with a filter afterwards.

WHAT THE CLASSES MEAN

stream is a mapped watercourse. short is a stub too brief to call one. The
unclassed remainder was drawn before the class field existed."""
    )

    spec[32] = dict(
        kicker=f"{f(pads['n'])} well pads, {f(pads['total_area_ha'], 1)} ha in total",
        bullets=[
            f"Average pad {f(pads['mean_area_m2'])} m2, "
            f"median {f(pads['median_area_m2'])} m2",
            f"Middle 80% between {f(pads['p10_area_m2'])} and "
            f"{f(pads['p90_area_m2'])} m2",
            area_line(pads, "area_ha"),
        ],
        notes=f"""HOW MANY PADS, AND HOW BIG

{f(pads['n'])} pads drawn, {f(pads['total_area_ha'], 1)} hectares of ground
between them.

The average is {f(pads['mean_area_m2'])} square metres. Picture a square about
{f(pads['mean_area_m2'] ** 0.5)} m on a side, which is roughly two tennis courts.

The average and the median ({f(pads['median_area_m2'])} m2) sit close together,
which means there is no long tail of monsters dragging the average up. Eight
pads in ten fall between {f(pads['p10_area_m2'])} and {f(pads['p90_area_m2'])} m2.

WHERE THEY ARE

{area_line(pads, 'area_ha')}

None on 613590. Pads were only ever drawn on 9t and on the eastern ground."""
    )

    spec[34] = dict(
        kicker=f"{f(out['n'])} pit outlines, {f(out['total_area_ha'], 2)} ha",
        bullets=[
            f"Average outline {f(out['mean_area_m2'])} m2, "
            f"median {f(out['median_area_m2'])} m2",
            f"Middle 80% between {f(out['p10_area_m2'])} and "
            f"{f(out['p90_area_m2'])} m2",
            area_line(out, "area_ha"),
            f"{f(wall['n'])} pair with a floor, giving "
            f"{f(wall['total_area_ha'], 2)} ha of measurable wall",
        ],
        notes=f"""THE OUTSIDE OF THE PIT

This is the big outline: the rim, the whole dent in the ground including its
sloping sides.

{f(out['n'])} of them drawn, {f(out['total_area_ha'], 2)} hectares in total.
Average {f(out['mean_area_m2'])} m2, median {f(out['median_area_m2'])} m2. A
typical one is a circle about {f(dia(out['median_area_m2']))} m across.

Eight in ten fall between {f(out['p10_area_m2'])} and {f(out['p90_area_m2'])} m2.

WHERE THEY ARE

{area_line(out, 'area_ha')}

AND THE WALL

Subtract the floor from the outline and what is left is the wall, the sloping
ring between the two. {f(wall['n'])} pits have both drawn, so only those
{f(wall['n'])} have a wall to measure: {f(wall['total_area_ha'], 2)} hectares of
it, averaging {f(wall['mean_area_m2'])} m2 per pit."""
    )

    ratio = 100.0 * ins["mean_area_m2"] / out["mean_area_m2"]
    spec[35] = dict(
        kicker=f"{f(ins['n'])} pit floors, {f(ins['total_area_ha'], 2)} ha",
        bullets=[
            f"Average floor {f(ins['mean_area_m2'])} m2, "
            f"median {f(ins['median_area_m2'])} m2",
            f"Middle 80% between {f(ins['p10_area_m2'])} and "
            f"{f(ins['p90_area_m2'])} m2",
            area_line(ins, "area_ha"),
            f"A floor is {ratio:.0f}% of its outline by area, the rest is wall",
        ],
        notes=f"""THE INSIDE OF THE PIT

This is the small blob in the middle: the flat floor at the bottom, not the
sloping sides.

{f(ins['n'])} drawn, {f(ins['total_area_ha'], 2)} hectares in total. That total
is small because each one is small. Average {f(ins['mean_area_m2'])} m2, median
{f(ins['median_area_m2'])} m2. A typical floor is a circle about
{f(dia(ins['median_area_m2']), 1)} m across, which is shorter than a parking
space is long.

Eight in ten fall between {f(ins['p10_area_m2'])} and {f(ins['p90_area_m2'])} m2.

WHY THAT MATTERS

The floor is only {ratio:.0f}% of the outline by area. Everything else is wall.
At half-metre pixels the average floor covers about
{f(ins['mean_area_m2'] / 0.25)} pixels, which is a small target. That is the
reason the pit model is scored on whether it finds the pit at all, rather than
on getting its edge exactly right.

WHERE THEY ARE

{area_line(ins, 'area_ha')}

The {ins['by_area']['613590']['n']} on 613590 are the ones the transfer test
scores against."""
    )
    return spec


def fill(tb, kicker, bullets):
    """Append the body under the title already sitting in this textbox."""
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].space_after = Pt(14)

    def para(text, size, bold, colour, space):
        p = tf.add_paragraph()
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Calibri"

    para(kicker, 18, True, INK, 12)
    for b in bullets:
        if b:
            para("•  " + b, 15, False, INK2, 9)


def set_notes(slide, text):
    """Write speaker notes even when the notes master has no body placeholder.

    Same fallback as _build_deck_v6.set_notes: python-pptx returns None from
    notes_text_frame when the master carries no body placeholder, and presenter
    view reads a type="body" placeholder on the notes slide.
    """
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is not None:
        tf.text = text
        return
    from lxml import etree
    from pptx.oxml.ns import qn
    tree = ns.shapes._spTree
    sp = etree.SubElement(tree, qn("p:sp"))
    nv = etree.SubElement(sp, qn("p:nvSpPr"))
    cnv = etree.SubElement(nv, qn("p:cNvPr"))
    cnv.set("id", "91")
    cnv.set("name", "Notes Placeholder")
    csp = etree.SubElement(nv, qn("p:cNvSpPr"))
    etree.SubElement(csp, qn("a:spLocks")).set("noGrp", "1")
    nvpr = etree.SubElement(nv, qn("p:nvPr"))
    etree.SubElement(nvpr, qn("p:ph")).set("type", "body")
    etree.SubElement(sp, qn("p:spPr"))
    tx = etree.SubElement(sp, qn("p:txBody"))
    etree.SubElement(tx, qn("a:bodyPr"))
    etree.SubElement(tx, qn("a:lstStyle"))
    for line in text.split("\n"):
        p = etree.SubElement(tx, qn("a:p"))
        r = etree.SubElement(p, qn("a:r"))
        etree.SubElement(r, qn("a:rPr")).set("lang", "en-US")
        etree.SubElement(r, qn("a:t")).text = line


def main() -> int:
    stats = json.loads(STATS.read_text(encoding="utf-8"))
    spec = build(stats)
    prs = Presentation(DECK)
    for n in SLIDES:
        d = spec[n]
        slide = prs.slides[n - 1]
        tbs = [sh for sh in slide.shapes if sh.has_text_frame]
        if not tbs:
            print(f"  slide {n}: no textbox, skipped")
            continue
        # the left column is the leftmost text shape; the figure has none
        tb = min(tbs, key=lambda sh: sh.left)
        title = tb.text_frame.paragraphs[0].runs[0].text
        if len(tb.text_frame.paragraphs) > 1:
            print(f"  slide {n}: body already present, skipped")
            continue
        fill(tb, d["kicker"], d["bullets"])
        set_notes(slide, d["notes"])
        print(f"  slide {n}  {title}")
        print(f"      {d['kicker']}")
    prs.save(DECK)
    print(f"\n  {DECK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

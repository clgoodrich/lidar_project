"""The rest of the further-research section: more ground, and a better model.

WHY
---
Retiring the old "Future Work" slide took out real ideas along with the stale
ones, and the NISAR slide only replaced one of them. This puts the rest back,
re-grounded in numbers from this repo rather than the retired pipeline's
metrics, and adds the two directions that were never on a slide at all --
other counties and older fields.

Everything asserted here is checked:

more ground
  McKean          PA Northcentral 2019 B19, QL1, already downloaded. The 18 deg
                  vendor scan-angle cut hits 6 of 7 Venango tiles and 0 of 7
                  McKean tiles (docs/analysis_log.md, "Not a convention"), so
                  the corduroy problem is a property of the March 2020 +/-20 deg
                  block, not of lidar.
  older fields    8,054 of 20,108 wells in the DEP Venango export carry
                  SPUD_DATE 1800-01-01, which is a placeholder, not a date.
                  Another 1,553 have none at all. So roughly 48% of the
                  county's recorded wells have no usable spud date -- and the
                  undated ones are the old shallow ones this method targets.
                  NOTE: this is deliberately NOT phrased as "8,054 wells
                  predate 1900". The sentinel year makes that reading available
                  and it would be false.
  Permian         TX_WestTexas_2018, QL1, 12-15 pts/m2 measured over the well
                  hotspots; RRC "Orphan Wells" layer gives independent truth,
                  and the best cell holds 136 orphans in 5 km
                  (docs/analysis_log.md, 2026-06-09).

a better model
  channels        notebooks/wellsight_v2/_dl.py DEFAULT_CHANNELS is exactly
                  seven bare-earth terrain layers. chm_9t_05.tif and
                  intensity_ground_9t_05.tif are both built and neither reaches
                  any detector.
  negatives       matches the standing preference for fixing confusion in
                  training rather than post-hoc filters.
  geomorphons     the old slide claimed +31% PR-AUC. That was measured on the
                  XGBoost/LightGBM ensemble which is no longer in the deck, so
                  the number is carried as "needs re-testing" rather than
                  quoted as if it still applied.

Placed before the NISAR slide so the section reads: more ground, a better
model, then watching what we found.

Run:
    python docs/presentation/_add_further_research_slides_v6.py
Writes (in place):
    docs/presentation/WellSight_Presentation v6.pptx
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs/presentation/WellSight_Presentation v6.pptx"
PML = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

INK = RGBColor(0x14, 0x1A, 0x1F)
INK2 = RGBColor(0x54, 0x5C, 0x63)
ACCENT = RGBColor(0x00, 0x96, 0xC7)

#: the new slides go immediately before this one
BEFORE_TITLE = "Further research — watching what we found"

SLIDES = [
    dict(
        title="Further research — more ground",
        kicker="Three ways to widen it, each with the data already downloaded",
        bullets=[
            "More of Pennsylvania — McKean is QL1, already held, and a "
            "cleaner delivery",
            "The 18° vendor cut hits 6 of 7 Venango tiles and 0 of 7 "
            "McKean tiles, so the corduroy is this flight block, not lidar",
            "Older fields — 8,054 of 20,108 Venango wells in the DEP "
            "export carry a placeholder spud date, not a real one",
            "Another 1,553 have none at all: the wells with the worst records "
            "are the old shallow ones this method is for",
            "A second basin — Permian QL1 at 12–15 pts/m², with "
            "Texas RRC orphan records as independent ground truth",
            "One Permian cell holds 136 confirmed orphans in 5 km, a denser "
            "test than anywhere in Venango",
        ],
        notes=(
            "Three directions, and none of them needs new data collection.\n\n"
            "McKean first, because we already have it and it is QL1 rather "
            "than QL2. It is also the cleaner delivery: the eighteen-degree "
            "vendor cut that causes our corduroy shows up in six of seven "
            "Venango tiles and zero of seven McKean tiles. That is worth "
            "saying out loud, because it means the artefact belongs to the "
            "March 2020 flight block, not to airborne lidar.\n\n"
            "Older fields is the interesting one. Of twenty thousand DEP "
            "records for Venango, eight thousand carry a spud date of "
            "January 1800, which is a placeholder the database uses when it "
            "does not know. Another fifteen hundred have nothing at all. Be "
            "careful how you say this: it does not mean eight thousand wells "
            "predate 1900. It means the records are missing, and missing "
            "records is exactly the population we are trying to find on the "
            "ground.\n\n"
            "The Permian is the generalisation test with real teeth: denser "
            "lidar, sparse vegetation, and the Texas Railroad Commission "
            "publishes a confirmed orphan layer, so for once there is ground "
            "truth we did not draw ourselves."
        ),
    ),
    dict(
        title="Further research — a better model",
        kicker="Carried over from the retired pipeline, and still worth doing",
        bullets=[
            "The detectors see seven bare-earth channels and nothing else: "
            "LRM, slope, TPI, openness, roughness",
            "Canopy height and return intensity are already built, and neither "
            "reaches any model",
            "A negative class for the confusers — ponds, cellars, quarry "
            "scrapes — trained in rather than filtered out afterwards",
            "Geomorphon enclosure as an extra channel; the old +31% was "
            "measured on the retired ensemble, so it needs re-testing",
            "Every one of these is a channel or a label change, not a new "
            "architecture",
        ],
        notes=(
            "This is the part of the old Future Work slide that survived, with "
            "the stale numbers taken off it.\n\n"
            "Start with what the model does not see. Seven channels, all "
            "derived from the bare-earth surface. We built a canopy height "
            "model and a ground-intensity raster, and neither one is wired "
            "into any detector. A pit under different canopy, or a pad with a "
            "different surface material, is invisible to the network as a "
            "material question -- it only ever sees shape.\n\n"
            "The negative class is the one I would do first. Ponds, cellars "
            "and quarry scrapes look like pits, and the right fix is to label "
            "them and train them in as negatives, not to filter them out after "
            "the fact.\n\n"
            "On geomorphons: the old slide claimed a thirty-one percent "
            "PR-AUC gain. That was measured on the gradient-boosted ensemble "
            "that is no longer in this deck, so treat it as untested against "
            "the U-Nets.\n\n"
            "The point of the last line is cost. None of this is a new "
            "architecture. It is extra channels and better labels, on the "
            "split we already have."
        ),
    ),
]


def rgb_runs(tf, runs):
    first = True
    for text, size, bold, colour, space in runs:
        if first and len(tf.paragraphs) and not tf.paragraphs[0].runs:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Calibri"


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def main() -> int:
    prs = Presentation(DECK)
    tgt_i = next(i for i, s in enumerate(prs.slides)
                 if title_of(s) == BEFORE_TITLE)
    src = prs.slides[tgt_i]
    sld_lst = prs.slides._sldIdLst

    for n, spec in enumerate(SLIDES):
        new = prs.slides.add_slide(src.slide_layout)
        for sh in list(new.shapes):
            sh._element.getparent().remove(sh._element)
        bg = src._element.find(f"{PML}cSld").find(f"{PML}bg")
        if bg is not None:
            new._element.find(f"{PML}cSld").insert(0, copy.deepcopy(bg))

        tb = new.shapes.add_textbox(Inches(1.00), Inches(0.62),
                                    Inches(11.5), Inches(0.9))
        # word_wrap off leaves wrap="none" + spAutoFit, which renderers
        # re-centre; the title then does not line up with its neighbours
        tb.text_frame.word_wrap = True
        rgb_runs(tb.text_frame, [(spec["title"], 30, True, ACCENT, 0)])

        body = new.shapes.add_textbox(Inches(1.00), Inches(1.78),
                                      Inches(11.2), Inches(4.6))
        body.text_frame.word_wrap = True
        runs = [(spec["kicker"], 18, True, INK, 14)]
        for b in spec["bullets"]:
            runs.append(("•  " + b, 16, False, INK2, 10))
        rgb_runs(body.text_frame, runs)

        foot = new.shapes.add_textbox(Inches(1.00), Inches(6.62),
                                      Inches(11.0), Inches(0.5))
        foot.text_frame.word_wrap = True
        rgb_runs(foot.text_frame,
                 [("Colton Goodrich  ·  University of Houston",
                   13, False, INK2, 0)])
        new.notes_slide.notes_text_frame.text = spec["notes"]

        # add_slide appends; walk it back to just before the NISAR slide,
        # keeping the two in the order they are declared
        node = list(sld_lst)[-1]
        sld_lst.remove(node)
        list(sld_lst)[tgt_i + n].addprevious(node)
        print(f"  added: {spec['title']}")

    prs.save(DECK)
    out = Presentation(DECK)
    print()
    for i, s in enumerate(out.slides, 1):
        if title_of(s).startswith("Further research"):
            print(f"  slide {i}: {title_of(s)}")
    print(f"  {len(out.slides)} slides")
    print(f"  {DECK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

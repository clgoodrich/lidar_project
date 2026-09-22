# -*- coding: utf-8 -*-
"""v15 -> v17: explain what a training patch is, in slide 57's speaker notes.

WHY v17 AND NOT v16
-------------------
A v16 already exists, built from v15 before the user edited v15 by hand. The
user's v15 is now newer than that v16, so v16 is a dead branch. Forking to v17
leaves the user's file untouched and does not reuse a version number that
already means something else.

WHAT GOES IN, AND WHY THERE
---------------------------
Slide 57 compares four architectures with "everything else held identical". The
patch geometry is part of that held-identical setup, and it appears nowhere in
the deck -- so the one number a listener can actually picture, 128 m of ground
per training example, is missing.

The jitter is the half worth saying out loud. Patches are centred on annotated
pits, because a patch dropped at random in this tile is empty forest. Centring
alone would teach "the answer is in the middle", so each centre is displaced by
up to 30 m first. The guard that goes with it -- a jittered patch is clipped to
the training blocks -- is what stops a shifted patch reaching across the line
into the held-out fold, and it is the first thing a careful listener will ask.

Source of the numbers: notebooks/wellsight_v2/s3_train/_pit_unet_cv5.py,
PATCH = 256 cells at 0.5 m and JITTER_M = 30.0.

SLIDE 54 IS LEFT ALONE, DELIBERATELY
------------------------------------
Slide 54's pasted notes look glued together when dumped to a terminal
("...not a custom design.Cross-entropy is..."), but the breaks are there as
\x0b soft line breaks, which is what PowerPoint writes on paste. They render
correctly. There is nothing to repair, and rewriting them would only risk the
user's own edit.

Run:
    python docs/presentation/_build_v17_patch_jitter_notes.py
Reads:  docs/presentation/WellSight_Presentation v15.pptx
Writes: docs/presentation/WellSight_Presentation v17.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
SRC = HERE / "WellSight_Presentation v15.pptx"
DST = HERE / "WellSight_Presentation v17.pptx"

SLIDE_57_TITLE = "Model building \u2014 does a bigger network help?"

PATCH_NOTE = (
    "If asked what a training patch is:\n\n"
    "The tile is too big to feed a network at once, so training happens on "
    "small square cut-outs. Each one is 256 pixels at 0.5 m, which is 128 m on "
    "the ground.\n\n"
    "Patches are centred on annotated pits, otherwise almost every patch would "
    "be empty forest.\n\n"
    "The jitter is the important half. Without it the pit sits dead centre in "
    "every single patch, and the network can learn \u201cthe answer is in the "
    "middle\u201d instead of what a pit looks like. So each patch centre is "
    "shifted by a random amount up to 30 m before it is cut. The pit lands "
    "somewhere different every time it is seen.\n\n"
    "One guard goes with it. A jittered patch is clipped to the training "
    "blocks, so a shifted patch cannot reach across the line into the held-out "
    "fold."
)


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def main() -> int:
    if not SRC.exists():
        raise SystemExit(f"missing: {SRC}")
    prs = Presentation(SRC)
    n_before = len(prs.slides)

    hits = [s for s in prs.slides if title_of(s).strip() == SLIDE_57_TITLE]
    if len(hits) != 1:
        raise SystemExit(f"expected one {SLIDE_57_TITLE!r}, found {len(hits)}")
    slide = hits[0]

    old = slide.notes_slide.notes_text_frame.text.rstrip()
    if "jitter" in old.lower():
        raise SystemExit("slide 57 already explains jitter; nothing to do")
    slide.notes_slide.notes_text_frame.text = old + "\n\n" + PATCH_NOTE
    new = slide.notes_slide.notes_text_frame.text
    print(f"  {SLIDE_57_TITLE}")
    print(f"  notes {len(old):5d} -> {len(new):5d} chars, "
          f"existing text kept verbatim")

    prs.save(DST)

    out = list(Presentation(DST).slides)
    if len(out) != n_before:
        raise SystemExit("slide count changed; that was not the intent")
    if "128 m on the ground" not in out[56].notes_slide.notes_text_frame.text:
        raise SystemExit("slide 57 is not where the note landed")
    print(f"\n  {len(out)} slides, unchanged")
    print(f"  slide 57 = {title_of(out[56])!r}  (note verified in place)")
    print(f"  {DST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

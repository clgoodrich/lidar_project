# -*- coding: utf-8 -*-
"""Re-embed slide 20's figure in v17 after the void contrast was raised.

WHY
---
python-pptx copies an image into the package when it is added, so regenerating
the PNG on disk does not change a deck that already holds it. Slide 20's left
panel was rebuilt -- the void class now separates by hue AND by contrast rather
than by a flat alpha-0.30 wash that went invisible at 52% coverage -- so the
copy inside v17 has to be replaced.

Position, size and every other shape on the slide are preserved. The figure
keeps its aspect ratio, which is unchanged, so the box is reused as it stands.

Edits v17 in place. No new version: the user asked for versioning to stop.

Run:
    python docs/presentation/_refresh_slide20_void_contrast_v17.py
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation

HERE = Path(__file__).resolve().parent
DECK = HERE / "WellSight_Presentation v17.pptx"
FIG = (HERE / "figures_30to45min/v6"
       / "dem_with_and_without_deleted_returns_9t.png")
TITLE = "Data QA — the surface, built both ways"


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def main() -> int:
    for p in (DECK, FIG):
        if not p.exists():
            raise SystemExit(f"missing: {p}")

    prs = Presentation(DECK)
    n_before = len(prs.slides)
    hits = [s for s in prs.slides if title_of(s).strip() == TITLE]
    if len(hits) != 1:
        raise SystemExit(f"expected one {TITLE!r}, found {len(hits)}")
    slide = hits[0]

    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if len(pics) != 1:
        raise SystemExit(f"slide has {len(pics)} pictures, expected 1")
    old = pics[0]
    old_sha = old.image.sha1[:12]
    left, top, width, height = old.left, old.top, old.width, old.height

    with Image.open(FIG) as im:
        w, h = im.size
    box_ar = width / height
    fig_ar = w / h
    if abs(box_ar - fig_ar) / fig_ar > 0.02:
        raise SystemExit(f"aspect changed ({fig_ar:.3f} vs box {box_ar:.3f}); "
                         "reposition deliberately rather than stretching")

    old._element.getparent().remove(old._element)
    new = slide.shapes.add_picture(str(FIG), left, top,
                                   width=width, height=height)
    print(f"  {TITLE}")
    print(f"  picture {old_sha} -> {new.image.sha1[:12]}  ({w}x{h} px)")
    print(f"  box unchanged: {width/914400:.2f} x {height/914400:.2f} in")

    prs.save(DECK)
    out = Presentation(DECK)
    if len(out.slides) != n_before:
        raise SystemExit("slide count changed; that was not the intent")
    print(f"  {len(out.slides)} slides, unchanged")
    print(f"  {DECK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

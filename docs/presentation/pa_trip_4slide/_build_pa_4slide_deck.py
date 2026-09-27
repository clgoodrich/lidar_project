"""Build the 4-slide Pennsylvania deck as a .pptx: the problem, the area, how
lidar sees a well site, and what we want to do there.

No results slides. The user asked for the problem and the area, not statistics.

Images (all already in the repo):
  figures/dep_wells_by_status_on_hillshade_9t.png           (_build_pa_area_figures.py)
  figures/pa_counties_venango_mckean_drake_well_locator.png (_build_pa_area_figures.py)
  ../v19_maps_and_plain_results/figures/rrim_zoom_sequence_tile_to_pit_9t_05.png

Colours: dark #17222B, light #F7F8F6, blue #1F5FA8, amber #D97706 / #E9A04A.
The figure palettes are validated in their own builders (all pairs, no red/green).
Fonts are Cambria (headings) and Calibri (body), which ship with Office.

Run (system Python, where python-pptx is installed):
  python docs/presentation/pa_trip_4slide/_build_pa_4slide_deck.py
Writes:
  docs/presentation/pa_trip_4slide/Orphaned_Wells_Pennsylvania_problem_and_area_4slides.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
V19 = HERE.parent / "v19_maps_and_plain_results" / "figures"
OUT = HERE / "Orphaned_Wells_Pennsylvania_problem_and_area_4slides.pptx"

DARK, LIGHT = RGBColor(0x17, 0x22, 0x2B), RGBColor(0xF7, 0xF8, 0xF6)
BLUE, AMBER_LT = RGBColor(0x1F, 0x5F, 0xA8), RGBColor(0xE9, 0xA0, 0x4A)
BODY_ON_LIGHT, MUTED_ON_LIGHT = RGBColor(0x3E, 0x4A, 0x55), RGBColor(0x56, 0x62, 0x6D)
BODY_ON_DARK, TITLE_ON_DARK = RGBColor(0xC9, 0xD3, 0xDA), RGBColor(0xF5, 0xF2, 0xEC)
CARD = RGBColor(0x22, 0x32, 0x40)
HEAD, TEXT = "Cambria", "Calibri"
M = 0.6  # side margin, inches


def background(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    """paras: list of (runs, size, space_after) where runs = [(str, color, bold, font)]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (runs, size, after) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        for s, color, bold, font in runs:
            r = p.add_run()
            r.text = s
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = font
    return tb


def eyebrow(slide, label, color):
    text(slide, M, 0.55, 8, 0.35, [([(label.upper(), color, True, TEXT)], 13, 0)])


def title(slide, s, color, y=0.9, h=0.8, w=12.1):
    text(slide, M, y, w, h, [([(s, color, True, HEAD)], 36, 0)])


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(w) if w else None, Inches(h) if h else None)


def body(s, color, size=18, after=10, lead=None, lead_color=None):
    runs = ([(lead + " ", lead_color, True, TEXT)] if lead else []) + [(s, color, False, TEXT)]
    return (runs, size, after)


def main() -> int:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. The problem ------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, LIGHT)
    eyebrow(s, "1 · The problem", BLUE)
    title(s, "The oil came before the paperwork", DARK, h=1.4, w=5.9)
    text(s, M, 2.55, 5.9, 3.6, [
        body("Wells were drilled here for decades before anyone had to record them.", BODY_ON_LIGHT, after=14),
        body("An abandoned, unplugged well can leak methane and brine. The open hole is a hazard too.",
             BODY_ON_LIGHT, after=14),
        body("Many old wells are on no list. Others are listed in the wrong place, or with no date.",
             BODY_ON_LIGHT, after=14),
    ])
    text(s, M, 6.45, 5.9, 0.5, [body("Right: one 4.5 km tile in Venango County, with every DEP well record on it.",
                                     MUTED_ON_LIGHT, size=12, after=0)])
    picture(s, FIG / "dep_wells_by_status_on_hillshade_9t.png", 6.85, 0.55, h=6.4)
    s.notes_slide.notes_text_frame.text = (
        "The problem is simple to state. Oil drilling here started in 1859, and for decades nobody had to "
        "record a well. When an old well is abandoned without a plug, it can leak methane into the air and "
        "brine into the ground and water. The open hole is also a physical hazard. The state keeps lists of "
        "orphaned and abandoned wells, but they are incomplete. Many old wells are on no list at all. Others "
        "are listed at the wrong spot, or with no drilling date. In Venango County about half of the 20,108 "
        "DEP records have no usable drilling date. The map is one 4.5 km square of our study area. Every mark "
        "is a DEP record: 171 orphaned or abandoned, 271 plugged and 616 active. The wells we care about most "
        "are the ones that are not on this map.")

    # 2. The area ---------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, LIGHT)
    eyebrow(s, "2 · The area", BLUE)
    title(s, "Venango County, where the oil industry began", DARK)
    picture(s, FIG / "pa_counties_venango_mckean_drake_well_locator.png", M, 1.95, w=8.2)
    text(s, 9.2, 2.2, 3.55, 4.6, [
        body("America's first commercial oil well is here.", BODY_ON_LIGHT, lead="Drake Well, 1859.",
             lead_color=DARK, after=18),
        body("Steep hills and deep stream valleys, under thick hardwood forest.", BODY_ON_LIGHT,
             lead="Hard ground to search.", lead_color=DARK, after=18),
        body("is old oil country too. We already hold its lidar.", BODY_ON_LIGHT, lead="McKean County",
             lead_color=DARK, after=0),
    ])
    s.notes_slide.notes_text_frame.text = (
        "Our study area is Venango County in northwest Pennsylvania. This is where the American oil industry "
        "started. Drake's well, drilled in 1859, sits at the north edge of the county, marked by the star. "
        "The boom that followed put thousands of wells into these hills. The ground itself makes the search "
        "hard. It is Appalachian Plateau: steep slopes, deep stream valleys and thick hardwood forest. From "
        "the air or in a satellite photo, the trees hide almost everything. McKean County, in light blue, is "
        "another old oil region to the northeast. We already have its lidar, so it is the natural next place "
        "to look.")

    # 3. How we look ------------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, DARK)
    eyebrow(s, "3 · How we look", AMBER_LT)
    title(s, "Old drilling still marks the forest floor", TITLE_ON_DARK)
    picture(s, V19 / "rrim_zoom_sequence_tile_to_pit_9t_05.png", M, 2.0, w=12.13)
    for i, (lead, rest) in enumerate([("Pit", "a dug hollow beside the wellhead"),
                                      ("Pad", "the levelled clearing the rig stood on"),
                                      ("Road", "the track that served the site")]):
        text(s, M + i * 4.17, 5.45, 3.8, 1.0,
             [([(lead + "  ", TITLE_ON_DARK, True, TEXT), (rest, BODY_ON_DARK, False, TEXT)], 18, 0)])
    s.notes_slide.notes_text_frame.text = (
        "Lidar is a laser flown over the ground. Enough pulses pass through gaps in the canopy to map the bare "
        "ground under the trees. This relief map is built from that ground surface. Hollows are blue and "
        "ridges are orange. Left to right, we zoom from the whole 4.5 km tile in Venango County to one pit "
        "10 m across. Every old well site leaves up to three marks: the pit dug beside the wellhead, the pad "
        "the rig stood on, and the road that served it. These marks outlast the wells and the paperwork.")

    # 4. In Pennsylvania --------------------------------------------------------
    s = prs.slides.add_slide(blank)
    background(s, DARK)
    eyebrow(s, "4 · In Pennsylvania", AMBER_LT)
    title(s, "What we want to do on the ground", TITLE_ON_DARK)
    cards = [
        ("Check candidates in person",
         "Visit a blind, random sample, strong scores and weak. This measures how many hits are real."),
        ("Look where no well is recorded",
         "Sites far from any DEP record may be undocumented wells. We report them as candidates for DEP to confirm."),
        ("Cover more ground",
         "Run the detectors over McKean County. Its denser 2019 lidar is already in hand."),
        ("Watch the sites over time",
         "Use NISAR radar from snow-free passes to ask whether the ground over a cluster of wells is moving."),
    ]
    cw, ch, gap = 5.9, 1.95, 0.33
    for i, (h, b) in enumerate(cards):
        x = M + (i % 2) * (cw + gap)
        y = 2.0 + (i // 2) * (ch + gap)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(cw), Inches(ch))
        box.adjustments[0] = 0.06
        box.fill.solid()
        box.fill.fore_color.rgb = CARD
        box.line.fill.background()
        box.shadow.inherit = False
        text(s, x + 0.35, y + 0.3, cw - 0.7, ch - 0.5, [
            ([(h, TITLE_ON_DARK, True, HEAD)], 22, 8),
            ([(b, BODY_ON_DARK, False, TEXT)], 16, 0),
        ])
    text(s, M, 6.75, 12.1, 0.4, [([("[Dates]  ·  [Number of sites]  ·  [Who we would work with]",
                                    RGBColor(0x9A, 0xA8, 0xB3), False, TEXT)], 13, 0)])
    s.notes_slide.notes_text_frame.text = (
        "This is what a trip to Pennsylvania is for. First, precision. Our annotations miss real pits, so only "
        "visiting sites tells us how many hits are real. The sample is blind and random, so we see weak scores "
        "as well as strong ones. Second, the real prize: sites that look like a well but sit far from any DEP "
        "record. Those are possible undocumented wells. We call them candidates until DEP confirms them. "
        "Third, scale. McKean County has denser 2019 lidar, already downloaded. Fourth, time. NISAR is a radar "
        "satellite that passes every 12 days. It cannot see a pit, but from snow-free passes it can tell "
        "whether the ground over a cluster of wells is sinking.")

    prs.save(OUT)
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

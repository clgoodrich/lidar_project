# -*- coding: utf-8 -*-
"""Build v19 from v18: more of the map, and a results section you can say out loud.

WHAT CHANGED, AND WHY (feedback after presenting v18, 2026-09-23)
-----------------------------------------------------------------
1. "Not enough emphasis on what we see on the maps." A five-slide map section
   goes in right after the RRIM slides: a zoom from the whole tile to one pit,
   then a pit, a pad and a road, then one well site with all three. RRIM only.
   No outlines, no arrows. The slide text says where to look.
2. "The results are over-complex." The v18 results slides put four metrics,
   two matching rules and a threshold rule on one caption line. v19 defines
   three words once (found, matched, search area) on a scoring slide, then
   gives every detector the same shape: one number per word, one plain
   takeaway, and the technical definition in one small line at the bottom.

v18 IS NOT EDITED. It is copied, and the copy is changed. Superseded results
slides are not deleted. They move to a "Backup" appendix at the end, so every
v18 number is still one click away during questions.

Every number here is one the 2026-09-22 audit already checked, re-derived
where it is new:
    pits   467 of 503 found, 738 flagged   pit_cv5_per_fold_9t.csv (F2 rows)
    pads   593 of 650 found, 1,010 flagged pad_cv5_per_fold_9t.csv (F2 rows)
    search area 0.21% / 9.8% / 5.02%       LEADERBOARD.md threshold-sweep table
    roads  722 of 735, 42.56 of 43.07 km    LEADERBOARD.md, v18 slide 57 notes
    613590 pits 139 of 153 (0.911), roads 0.759 / 0.828   v18 slides 65-66
    architectures 0.559/0.561, 0.555/0.608  arch_compare_summary_9t_1m.csv

Run with the system Python (python-pptx lives there, not in .venv):
    python docs/presentation/v19_maps_and_plain_results/_build_v19_from_v18.py
Writes:
    docs/presentation/v19_maps_and_plain_results/WellSight_Presentation v19.pptx
"""
from __future__ import annotations

import copy
import io
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn
from PIL import Image

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / "WellSight_Presentation v18.pptx"
DST = HERE / "WellSight_Presentation v19.pptx"
FIG = HERE / "figures"

BG = RGBColor(0xF7, 0xF8, 0xF6)
TITLE = RGBColor(0x00, 0x96, 0xC7)
INK = RGBColor(0x14, 0x1A, 0x1F)
INK2 = RGBColor(0x54, 0x5C, 0x63)
ACCENT = RGBColor(0x1F, 0x5F, 0xA8)
RULE = RGBColor(0xDD, 0xE0, 0xDA)
FONT = "Calibri"


# ---------------------------------------------------------------- helpers
def new_slide(prs):
    layout = next(l for l in prs.slide_layouts if l.name == "Blank")
    s = prs.slides.add_slide(layout)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(s, x, y, w, h, lines, size=16, color=INK2, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    """lines: str, or list of str / (str, dict) where dict overrides size/color/bold."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        opt = {}
        if isinstance(ln, tuple):
            ln, opt = ln
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(opt.get("space_after", space_after))
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.color.rgb = opt.get("color", color)
    return tb


def title(s, t, sub=None, w=12.2, tw=None):
    """tw: title box width (defaults to w). w: subtitle box width."""
    text(s, 0.55, 0.30, tw or w, 0.7, t, size=28, color=TITLE, bold=True)
    if sub:
        text(s, 0.58, 1.02, w, 0.5, sub, size=18, color=INK, bold=True)


def bullets(s, x, y, w, h, items, size=16):
    """Real bullets (a:buChar) with a hanging indent, so wrapped lines align."""
    tb = text(s, x, y, w, h, list(items), size=size, color=INK2, space_after=9)
    for para in tb.text_frame.paragraphs:
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.26))))
        pPr.set("indent", str(-int(Inches(0.26))))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
    return tb


def picture(s, src, x, y, w, h, align="center", valign="middle"):
    """Fit an image (path or bytes) inside the box, keeping its aspect."""
    if isinstance(src, (bytes, bytearray)):
        im = Image.open(io.BytesIO(src)); stream = io.BytesIO(src)
    else:
        im = Image.open(src); stream = str(src)
    iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else (x if align == "left" else x + w - pw)
    py = y if valign == "top" else y + (h - ph) / 2
    s.shapes.add_picture(stream, Inches(px), Inches(py), Inches(pw), Inches(ph))
    return px, py, pw, ph


def stat(s, x, y, w, number, label, num_size=40):
    text(s, x, y, w, 0.7, number, size=num_size, color=ACCENT, bold=True)
    text(s, x, y + 0.72, w, 0.6, label, size=15, color=INK2)


def notes(s, body):
    s.notes_slide.notes_text_frame.text = body


def largest_picture_blob(slide):
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    return max(pics, key=lambda p: p.width * p.height).image.blob


# ---------------------------------------------------------------- map section
def map_slides(prs):
    out = []

    # M1 zoom sequence
    s = new_slide(prs); out.append(s)
    title(s, "Zooming in on the ground",
          "The whole 9t tile down to a single pit, in four steps.")
    x0, y0, W = 0.55, 1.85, 12.23
    _, _, pw, ph = picture(s, FIG / "rrim_zoom_sequence_tile_to_pit_9t_05.png",
                           x0, y0, W, 3.2)
    panel = pw * 700 / 2872; gap = pw * 24 / 2872
    caps = ["The whole tile, 4.5 km", "1 km", "250 m", "50 m. One pit"]
    for i, c in enumerate(caps):
        text(s, x0 + i * (panel + gap), y0 + ph + 0.12, panel, 0.4, c,
             size=15, color=INK, bold=True, align=PP_ALIGN.CENTER)
    bullets(s, 0.55, y0 + ph + 0.75, 12.2, 1.4, [
        "The white box in each view is the next view.",
        "Teal is ground that dips. Bright is ground that bulges. Red is steep.",
        "Everything the models look for is in the last two panels. It is small.",
    ])
    notes(s, "Four views of the same ground, each one a zoom into the box on the "
             "one before.\n\nLeft is the whole 9t tile, 4.5 km on a side. The "
             "branching teal lines are stream valleys.\n\nThe last panel is 50 m "
             "across and holds one pit. The pit is the bright ring with a teal "
             "center.\n\nHow to read the colors, once, because every map from "
             "here uses them. Teal means the ground curves down, like a valley or "
             "a pit. Bright yellow-white means it curves up, like a rim or a "
             "ridge. Gray is flat. Red means steep.\n\nThe point of the slide is "
             "scale. A pit is 15 m across on a tile 4,500 m across. That is why "
             "nobody finds these by looking at the whole map.")

    # M2 pits
    s = new_slide(prs); out.append(s)
    title(s, "What a pit looks like", "A teal floor inside a bright rim.", w=5.7)
    bullets(s, 0.55, 1.65, 5.7, 2.1, [
        "About 15 m across the rim. The floor is about 6 m.",
        "About half a meter deep. It is a dish, not a hole.",
        "Teal marks the hollow. The bright ring is the raised rim around it.",
    ])
    picture(s, FIG / "rrim_pit_deeper_50m_9t_05.png", 0.55, 3.75, 2.6, 2.6)
    picture(s, FIG / "rrim_pit_shallower_50m_9t_05.png", 3.55, 3.75, 2.6, 2.6)
    text(s, 0.55, 6.42, 2.6, 0.5, "A deeper pit (90th percentile)", size=12,
         align=PP_ALIGN.CENTER)
    text(s, 3.55, 6.42, 2.6, 0.5, "A shallower pit (20th percentile)", size=12,
         align=PP_ALIGN.CENTER)
    picture(s, FIG / "rrim_pit_typical_50m_9t_05.png", 6.55, 0.45, 6.3, 6.3)
    text(s, 6.55, 6.82, 6.3, 0.4, "A typical pit. Median size and depth. Each view is 50 m across.",
         size=12, align=PP_ALIGN.CENTER)
    notes(s, "Three real pits from 9t, with nothing drawn on them. This is what "
             "the eye has to find.\n\nThe big one is the typical pit. It was not "
             "picked by eye. It is the pit closest to the median in both size and "
             "depth, out of the paired pits in 9t.\n\nThe two small ones are a "
             "deeper pit, at the 90th percentile of depth, and a shallower one, "
             "at the 20th.\n\nWhere to look. In the big one the pit is dead "
             "center: a teal floor with a bright ring around it. The broad teal "
             "band at the bottom left is a road.\n\nWorth saying out loud. Depth "
             "is not the same as visibility. The shallow pit on the right is the "
             "easiest of the three to see, because it sits on flat ground with "
             "nothing else around it.\n\nThe numbers on the slide come from the "
             "morphology slide later on: 15.3 m across the rim, 5.8 m across the "
             "floor, 0.54 m deep at the median, measured on the 503 paired pits "
             "in 9t.")

    # M3 pads
    s = new_slide(prs); out.append(s)
    title(s, "What a pad looks like", "A flat, smooth patch cut into the hillside.", w=5.7)
    bullets(s, 0.55, 1.65, 5.7, 2.1, [
        "About 40 m across. Room for a rig, a tank and a truck.",
        "Flat ground shows gray. The cut edge where it meets the slope shows red.",
        "Harder to see than a pit. A flat patch is not unusual on its own.",
    ])
    picture(s, FIG / "rrim_pad_large_90m_9t_05.png", 0.55, 3.75, 2.6, 2.6)
    picture(s, FIG / "rrim_pad_small_90m_9t_05.png", 3.55, 3.75, 2.6, 2.6)
    text(s, 0.55, 6.42, 2.6, 0.5, "A larger pad (85th percentile)", size=12,
         align=PP_ALIGN.CENTER)
    text(s, 3.55, 6.42, 2.6, 0.5, "A smaller pad (20th percentile)", size=12,
         align=PP_ALIGN.CENTER)
    picture(s, FIG / "rrim_pad_typical_90m_9t_05.png", 6.55, 0.45, 6.3, 6.3)
    text(s, 6.55, 6.82, 6.3, 0.4, "A typical pad. Median area. Each view is 90 m across.",
         size=12, align=PP_ALIGN.CENTER)
    notes(s, "Three pads, again with nothing drawn on them.\n\nA pad is the flat "
             "area cleared for the drilling equipment. Picked by area: the big "
             "one is the median pad, the small ones are at the 85th and 20th "
             "percentiles. They were picked from pads with no pit drawn on them, "
             "so the pad would read on its own. Look closely and all three still "
             "show a small ring. Those are very likely pits nobody drew. That is "
             "worth saying: it is the same thing that makes the pad and pit "
             "precision numbers conservative later on.\n\nWhere to look in the "
             "big one. The pad is the smooth gray ground in the center and to the "
             "right. Its left edge is the sharp line where it was cut into the "
             "slope.\n\nIn the small images the pad is easier: a flat bench off a "
             "road, with a ring on it.\n\nSay this plainly, because it "
             "sets up the results. A pit is a shape that nature rarely makes. A "
             "flat patch is not. That is why the pad model flags more ground than "
             "the pit model, and it is visible here before any model runs.\n\n"
             "About 40 m across comes from the annotation slide: median pad in "
             "9t is 1,645 square meters.")

    # M4 roads
    s = new_slide(prs); out.append(s)
    title(s, "What a road looks like", "A narrow teal groove with a bright edge beside it.", w=5.7)
    bullets(s, 0.55, 1.65, 5.7, 2.1, [
        "Old access tracks, a few meters wide, worn into the ground.",
        "Bold roads are easy. Faint roads are most of the problem.",
        "Roads connect. A road that ends at a flat patch is how a well site shows itself.",
    ])
    picture(s, FIG / "rrim_road_bold_120m_9t_05.png", 0.55, 3.75, 2.6, 2.6)
    picture(s, FIG / "rrim_road_faint_120m_9t_05.png", 3.55, 3.75, 2.6, 2.6)
    text(s, 0.55, 6.42, 2.6, 0.5, "A bold road, 120 m view", size=12,
         align=PP_ALIGN.CENTER)
    text(s, 3.55, 6.42, 2.6, 0.5, "A faint road, 120 m view", size=12,
         align=PP_ALIGN.CENTER)
    picture(s, FIG / "rrim_roads_context_300m_9t_05.png", 6.55, 0.45, 6.3, 6.3)
    text(s, 6.55, 6.82, 6.3, 0.4, "A 300 m view with a typical amount of road for 9t.",
         size=12, align=PP_ALIGN.CENTER)
    notes(s, "Roads, with nothing drawn on them.\n\nThe big view is 300 m across. "
             "It was not picked for looks. The tile was cut into a 15 by 15 grid "
             "and this is the window at the 75th percentile of road length, so a "
             "busier-than-average piece of ground, not the busiest.\n\nThe roads "
             "are the thin teal lines. Teal means the ground dips, and a road is "
             "worn slightly below what is around it. Several run top to bottom and "
             "branch.\n\nThe two small views are from the bold-road and "
             "faint-road examples drawn by hand. Each is the median-length one of "
             "its kind.\n\nThe bold road on the left is obvious.\n\nThe faint one "
             "on the right is not, and that is the point. It climbs from the lower "
             "left to the center of the view and turns back south. Most people "
             "cannot see it at first. The strong teal line top right is a "
             "different, bolder feature.\n\nFaint roads are where the road model "
             "misses, and where the labels were thinnest.")

    # M5 well site
    s = new_slide(prs); out.append(s)
    title(s, "One well site, all three together", "A road, a pad beside it, and a pit on the pad.", w=5.7)
    bullets(s, 0.55, 1.65, 5.7, 3.4, [
        "The road is the teal curve running from the top left to the bottom right.",
        "The pad is the rounded flat loop off its left side.",
        "The pit is the bright ring just below the center, at the edge of the pad.",
        "This is the pattern the models learn. Roads lead to pads. Pads hold pits.",
    ])
    text(s, 0.55, 4.55, 5.7, 1.2,
         "Nothing is drawn on the map. Everything described here is in the ground.",
         size=15, color=INK, bold=True)
    picture(s, FIG / "rrim_well_site_road_pad_pit_150m_9t_05.png", 6.55, 0.45, 6.3, 6.3)
    text(s, 6.55, 6.82, 6.3, 0.4, "150 m view. The median pad among those holding a pit and a road.",
         size=12, align=PP_ALIGN.CENTER)
    notes(s, "One real well site, all three features in one view.\n\nPicked by "
             "rule: of the 229 pads in 9t that hold at least one pit and touch at "
             "least one road, this one has the median area.\n\nWalk the audience "
             "through it. The road comes in at the top and curves down to the "
             "bottom right. The pad is the loop on its left, flat and gray. The "
             "pit is the bright ring with a teal center just below the middle of "
             "the image, at the lower edge of the pad.\n\nThat sequence, road then "
             "pad then pit, is the whole idea. A person finds wells this way, and "
             "so do the models. They are trained on each feature separately, but "
             "the features come together on the ground.")
    return out


# ---------------------------------------------------------------- results section
HOW = ("How it was measured: five-fold spatial cross-validation on 9t. "
       "Found means the flagged shape overlaps the drawn one by at least 30% "
       "(IoU ≥ 0.30). The cut-off is set per fold to favor finding over "
       "flagging (F2).")


def results_slides(prs, v18):
    out = {}

    # R0 how we score
    s = new_slide(prs); out["score"] = s
    title(s, "How we score a detector",
          "Three questions. Every result that follows answers them.")
    rows = [
        ("Found", "Of the features drawn by hand, how many did the model flag?",
         "This is recall."),
        ("Matched", "Of everything the model flagged, how much sits on a drawn feature?",
         "This is precision. An extra flag may still be real, just never drawn."),
        ("Search area", "How much of the map would a person have to check?",
         "Flagged area as a share of the whole tile."),
    ]
    y = 1.75
    for word, q, d in rows:
        text(s, 0.55, y, 1.9, 0.5, word, size=20, color=ACCENT, bold=True)
        text(s, 2.45, y + 0.02, 3.6, 0.9, [(q, {"size": 15, "color": INK, "bold": True}),
                                            (d, {"size": 13, "color": INK2})], space_after=3)
        y += 1.3
    text(s, 0.55, 5.8, 5.6, 1.2,
         "Every score is on ground the model never trained on. The tile is cut "
         "into 144 blocks in five groups. Each model trains on three groups, "
         "tunes on one and is scored on the last.", size=13, color=INK2)
    picture(s, FIG / "scoring_found_missed_extra_schematic.png", 6.4, 1.6, 6.5, 4.4)
    notes(s, "Before any number, the three words the numbers mean.\n\nFound. Of "
             "the pits someone drew by hand, what share did the model flag. "
             "Technically that is recall.\n\nMatched. Of everything the model "
             "flagged, what share sits on a drawn feature. Technically that is "
             "precision. One caution belongs here, not later. A flag with nothing "
             "drawn under it is not necessarily wrong. It may be a real pit nobody "
             "drew. So matched is a floor on how often the model is right.\n\n"
             "Search area. How much of the map a person would have to look at. "
             "That is the cost of using the output.\n\nThe diagram. Dashed rings "
             "are the answer key. Blue blobs inside a ring are found. Red rings "
             "with nothing inside are missed. Orange blobs with no ring are extra "
             "flags.\n\nThe last line on the slide is the fairness rule. Nothing "
             "is scored on ground the model trained on. The tile is cut into 144 "
             "blocks, the blocks into five groups. Each of five models trains on "
             "three groups, uses one to pick its cut-off, and is scored on the "
             "last. Every pit gets scored exactly once, by a model that never saw "
             "it.\n\nIf asked what counts as found: the flagged shape has to "
             "overlap the drawn shape by at least 30%, measured as intersection "
             "over union.")

    def result(key, ttl, sub, stats, takeaway, how, img, img_note):
        s = new_slide(prs); out[key] = s
        title(s, ttl, sub, w=5.4, tw=12.2)
        y = 1.7
        for num, lab in stats:
            stat(s, 0.55, y, 5.0, num, lab)
            y += 1.35
        text(s, 0.55, y + 0.02, 5.0, 0.5, takeaway, size=16, color=INK, bold=True)
        text(s, 0.55, y + 0.62, 5.0, 0.8, how, size=11, color=INK2)
        px, py, pw, ph = picture(s, img, 5.75, 1.7, 7.3, 5.2, valign="top")
        text(s, px, py + ph + 0.1, pw, 0.4, img_note, size=11, align=PP_ALIGN.CENTER)
        return s

    s = result("pits", "Pits", "Nearly every pit found, and a short list.",
               [("467 of 503", "drawn pits found. 93%."),
                ("63%", "of what it flagged sits on a drawn pit. 467 of 738 flags."),
                ("0.21%", "of the map to search. A separate test found 126 of 127 pit rims.")],
               "The strongest result in the talk.", HOW,
               largest_picture_blob(v18[60]),
               "Pit probability beside the terrain, on 9t.")
    notes(s, "Pits, in the three words.\n\nFound: 467 of the 503 drawn pits. "
             "93%.\n\nMatched: it flagged 738 shapes, and 467 of them sit on a "
             "drawn pit. 63%. The other 271 are extras. Some of those are likely "
             "real pits nobody drew. When more pits were annotated between runs, "
             "this number went up, which is what you would expect if that is "
             "happening. It is a likely reading, not a proven one.\n\nSearch area: "
             "0.21% of the map. Be careful here. This number is from a SEPARATE "
             "test, on 127 withheld pit rims, at a fixed cut-off of 0.20, where a "
             "pit counts as found if the center of the flag falls inside the rim. "
             "It found 126 of the 127. Different test, different pits, so do not "
             "say it in the same breath as the 93%.\n\nWhat 0.21% means in "
             "practice: on a 20 square kilometer tile, about 4 hectares to walk.\n\n"
             "The image: terrain on the left, the model's pit probability on the "
             "right, same ground. Each bright dot on the right is a place the "
             "model thinks holds a pit floor. The full-slide version with the "
             "hand-drawn pits outlined is in the backup.\n\nThe small print on the slide, for a technical question: "
             "five-fold spatial cross-validation, IoU of at least 0.30 to count as "
             "found, and a cut-off chosen per fold by F2, which weights finding "
             "twice as heavily as flagging.")

    s = result("pads", "Pads", "It finds pads, but it flags too much ground.",
               [("593 of 650", "drawn pads found. 91%."),
                ("59%", "of what it flagged sits on a drawn pad. 593 of 1,010 flags."),
                ("9.8%", "of the map to search. A separate sweep found 178 of 194 pads.")],
               "Finding is fine. The search area is the problem.", HOW,
               largest_picture_blob(v18[59]),
               "Pad probability beside the terrain, on 9t.")
    notes(s, "Pads, the same three words.\n\nFound: 593 of 650 drawn pads. 91%. "
             "That is about as good as the pits.\n\nMatched: 1,010 flags, 593 on "
             "a drawn pad. 59%. Four in ten flags have nothing drawn under them.\n\n"
             "Search area: 9.8% of the map, from a separate sweep on withheld "
             "pads, where it found 178 of 194. That is a tenth of the tile to "
             "check. Pits needed a fifth of one percent.\n\nWhy. The map slides "
             "showed it. A pit is a shape nature rarely makes. A flat patch of "
             "ground is common. So the pad model has more look-alikes to reject.\n\n"
             "Quote 9.8%, not 11.6%. An earlier cut-off flagged 11.6% and found "
             "the same 178 pads. The higher cut-off gets the same finds with less "
             "ground.\n\nSame measurement rules as the pits: five-fold spatial "
             "cross-validation, IoU of at least 0.30, cut-off chosen by F2.")

    s = result("roads", "Roads", "The easiest target. Long and connected.",
               [("722 of 735", "withheld road segments found. 98%."),
                ("42.6 of 43.1 km", "of withheld road length recovered. 98.8%."),
                ("5.0%", "of the map flagged.")],
               "A road missed in one spot shows up in the next.",
               "How it was measured: roads are cut into segments of about 40 m. "
               "Only segments whose whole road was held out of training are "
               "scored. Cut-off 0.20.",
               largest_picture_blob(v18[57]),
               "Road probability beside the terrain, on 9t.")
    notes(s, "Roads.\n\nFound: roads are long, so they are cut into segments of "
             "about 40 m for scoring. 722 of 735 segments found, 98%. Those 735 "
             "are the segments whose entire road was held out of training, which "
             "is the strict version. A road half in training and half out would "
             "be too easy.\n\nBy length: 42.6 of 43.1 km recovered, 98.8%. The "
             "segment count and the length are two different measures and they "
             "agree.\n\nSearch area: 5.0% of the map.\n\nThere is no matched "
             "figure on this slide. Roads were scored for how much of the drawn "
             "road was found, not for how much of the flagged area was real. The "
             "second tile has that number, a couple of slides on.\n\nWhy roads "
             "are easiest: a road is a long, connected line. Miss a short stretch "
             "and the rest still gives it away.")

    # R5 summary on 9t
    s = new_slide(prs); out["summary"] = s
    title(s, "The three detectors, side by side",
          "Same tile, same rules, all scored on ground the model never trained on.")
    data = [["", "Found", "Matched", "Search area", "In one line"],
            ["Pits", "93%  (467 of 503)", "63%", "0.21%", "A short list"],
            ["Pads", "91%  (593 of 650)", "59%", "9.8%", "Too much ground"],
            ["Roads", "98%  (722 of 735)", "not scored this way", "5.0%", "The easiest target"]]
    shp = s.shapes.add_table(4, 5, Inches(0.55), Inches(1.85), Inches(12.2), Inches(3.0))
    tbl = shp.table
    widths = [1.6, 3.0, 2.5, 2.0, 3.1]
    for i, wd in enumerate(widths):
        tbl.columns[i].width = Inches(wd)
    for r in range(4):
        tbl.rows[r].height = Inches(0.75)
        for c in range(5):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xE9, 0xEC, 0xE6) if r == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.text = ""
            p = tf.paragraphs[0]
            run = p.add_run(); run.text = data[r][c]
            f = run.font; f.name = FONT
            f.size = Pt(18 if r == 0 or c == 0 else 18)
            f.bold = r == 0 or c == 0
            f.color.rgb = INK if (r == 0 or c in (0, 4)) else (ACCENT if c in (1, 3) else INK2)
            if r > 0 and c in (1, 3):
                f.bold = True
    text(s, 0.55, 5.2, 12.2, 1.3, [
        "Found and matched come from five-fold cross-validation. Search area comes "
        "from a separate threshold sweep on withheld ground. The two are different "
        "tests, so read each column on its own.",
        "Roads were scored by length recovered, not by how much of the flagged area "
        "was real. That second number exists for the new tile.",
    ], size=13, color=INK2)
    notes(s, "Everything so far, in one table.\n\nRead it by column. Found: all "
             "three are above 90%. Finding is not the problem for any of them.\n\n"
             "Search area is where they separate. Pits 0.21%. Roads 5%. Pads "
             "nearly 10%.\n\nThat is the one-line verdict for each. The pit model "
             "gives you a short list. The road model finds nearly all of it. The "
             "pad model finds pads but hands you a tenth of the map.\n\nThe small "
             "print is there because someone will ask. The found and matched "
             "columns and the search-area column are two different tests on two "
             "different withheld sets. They agree in direction but should not be "
             "multiplied together.")

    # R-arch: bigger network
    s = new_slide(prs); out["arch"] = s
    title(s, "Would a bigger network do better?", "For pits, no. For pads, maybe.", w=5.4, tw=12.2)
    bullets(s, 0.55, 1.65, 5.3, 3.4, [
        "Same data, same training. Only the network changed.",
        "Pits: 0.559 to 0.561 with three times the parameters. The folds vary by 0.020, so that is no gain.",
        "Pads: 0.555 to 0.608. Bigger than the noise. Worth another look.",
        "For pits the limit is the data and the labels, not the network.",
    ])
    text(s, 0.55, 5.6, 5.3, 1.3,
         "Score: overlap between predicted and drawn area (IoU) at 1 m, averaged "
         "over five folds. It is a different measure from the found rates on the "
         "results slides.", size=11, color=INK2)
    cd = CategoryChartData()
    cd.categories = ["Pits", "Pads"]
    cd.add_series("Plain U-Net, 7.8 M parameters", (0.559, 0.555))
    cd.add_series("Largest network tried, 26.1 M parameters", (0.561, 0.608))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.2), Inches(1.2),
                            Inches(6.7), Inches(5.6), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(13); ch.legend.font.name = FONT
    ch.legend.font.color.rgb = INK2
    va = ch.value_axis
    va.minimum_scale = 0.0; va.maximum_scale = 0.7
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RULE
    va.tick_labels.font.size = Pt(12); va.tick_labels.font.color.rgb = INK2
    va.format.line.fill.background()
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(16); ca.tick_labels.font.bold = True
    ca.tick_labels.font.color.rgb = INK
    plot = ch.plots[0]
    plot.gap_width = 70; plot.overlap = -8
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = "0.000"; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(14); dl.font.bold = True; dl.font.color.rgb = INK
    # Ordered pair, small -> large network, one hue light -> dark.
    # dataviz validate_palette.js "#4E8FE6,#1F5FA8" --pairs all --mode light
    # --surface #F7F8F6: all checks PASS, worst pair dE 16.3 deutan, 16.3 normal,
    # both >= 3:1. (Grey #8E959B was tried first and FAILED the chroma floor.)
    for ser, col in zip(plot.series, (RGBColor(0x4E, 0x8F, 0xE6), ACCENT)):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
    notes(s, "Whether a larger network would do better.\n\nFour networks were "
             "tried, from a plain U-Net to a U-Net++ on a pretrained ResNet-34. "
             "The chart shows the smallest and the largest. Everything else was "
             "held fixed: same folds, same seven channels, same loss, same 40 "
             "epochs.\n\nPits: 0.559 against 0.561. The five folds differ from "
             "each other by 0.020, so a gain of 0.002 is inside the noise. Three "
             "times the parameters bought nothing.\n\nPads: 0.555 against 0.608. "
             "That is bigger than the fold-to-fold spread, and every fold of the "
             "big network beats all but one fold of the small one. Suggestive. "
             "Needs more runs before it is a finding.\n\nThe score here is "
             "different from the found rates. It is the overlap between the "
             "predicted area and the drawn area, pixel by pixel, at 1 m. Do not "
             "compare it to the 93%.\n\nThe takeaway: for pits, more labels and "
             "better data will do more than a bigger network.")

    # R6 second tile
    s = new_slide(prs); out["tile2"] = s
    title(s, "On a tile it never saw", "613590. The same models, nothing retuned.", w=5.3)
    stat(s, 0.55, 1.75, 5.0, "139 of 153", "drawn pits found. 91%, against 93% at home.")
    stat(s, 0.55, 3.17, 5.0, "76%", "of the drawn road length covered.")
    stat(s, 0.55, 4.59, 5.0, "83%", "of the road it drew is real road.")
    text(s, 0.55, 5.95, 5.0, 0.5, "Moving to new ground cost the pit model two points.",
         size=16, color=INK, bold=True)
    text(s, 0.55, 6.5, 5.0, 0.9,
         "How it was measured: the five 9t models, thresholds frozen. Pits are the "
         "average of the five (range 127 to 146). Roads at cut-off 0.50 on the "
         "hand-drawn subset. Pads and drainage were never drawn here, so they are "
         "not scored.", size=11, color=INK2)
    # v18 slide 66's map has "146 found, recall 0.954" baked into it, which
    # contradicts the audited five-model average on this slide (139, 0.911).
    # Slide 63's map carries no numbers, so it is used instead.
    px, py, pw, ph = picture(s, largest_picture_blob(v18[63]), 5.75, 1.7, 7.3, 5.2,
                             valign="top")
    text(s, px, py + ph + 0.1, pw, 0.4, "Pit probability beside the terrain, on 613590.",
         size=11, align=PP_ALIGN.CENTER)
    notes(s, "The second tile, 613590. No model ever saw any of it.\n\nPits: on "
             "average the five models found 139 of the 153 drawn pits, 91%. On "
             "the home tile it was 93%. The individual models range from 127 to "
             "146. A model that had memorised the home tile would collapse here. "
             "This one drops two points.\n\nNothing was retuned. The cut-offs are "
             "the ones chosen on 9t. Retuning here would be fitting to the test.\n\n"
             "No matched figure for pits, because this tile is not fully drawn. "
             "An extra flag here could be a pit nobody got to.\n\nRoads: 76% of "
             "the drawn road length is covered, and 83% of the road the model drew "
             "is real road. That second number is the road version of matched.\n\n"
             "Pads and drainage were never drawn on this tile, so there is nothing "
             "to score them against.\n\nOne honest limit. Same county, same 2019 "
             "survey, same flight block. This shows the model moves to new ground. "
             "It does not show it moves to a new survey.\n\nIf asked about the "
             "second threshold rule: under the balanced F1 rule the drop is "
             "larger, 0.861 to 0.792. The backup slides have it.")

    # conclusion
    s = new_slide(prs); out["numbers"] = s
    title(s, "What the numbers say", "Three detectors, all scored on ground they never trained on.")
    rows = [("Pits", "93%", "found", "0.21%", "of the map to search"),
            ("Roads", "98%", "found", "5.0%", "of the map to search"),
            ("Pads", "91%", "found", "9.8%", "of the map to search")]
    y = 1.85
    for name, a, al, b, bl in rows:
        text(s, 0.9, y + 0.12, 2.0, 0.6, name, size=24, color=INK, bold=True)
        text(s, 3.0, y, 1.9, 0.8, a, size=40, color=ACCENT, bold=True)
        text(s, 4.95, y + 0.22, 1.6, 0.5, al, size=16, color=INK2)
        text(s, 6.9, y, 2.2, 0.8, b, size=40, color=ACCENT, bold=True)
        text(s, 9.15, y + 0.22, 3.6, 0.5, bl, size=16, color=INK2)
        y += 1.05
    text(s, 0.9, 5.2, 11.6, 1.5, [
        ("On a second tile the pit model found 91%, down from 93%.", {"size": 18, "color": INK, "bold": True}),
        ("Every number is against hand-drawn features, not field visits. Nothing here is a confirmed well.",
         {"size": 15, "color": INK2}),
    ], space_after=8)
    notes(s, "The whole results section in three lines.\n\nPits: 93% found, and a "
             "short list, a fifth of a percent of the map.\n\nRoads: 98% found, "
             "5% of the map.\n\nPads: 91% found, but nearly 10% of the map to "
             "check. That is the weak one.\n\nOn a tile no model had seen, pit "
             "finding dropped two points.\n\nAnd the limit, before anyone asks. "
             "All of this is measured against features drawn by hand on the map. "
             "Nobody has walked to these sites. Every detection is a candidate.\n\n"
             "The search-area numbers come from a separate threshold sweep, not "
             "from the same run as the found rates. They agree in direction.")

    # backup divider
    s = new_slide(prs); out["backup"] = s
    text(s, 0.9, 2.6, 11.5, 1.0, "Backup", size=40, color=TITLE, bold=True)
    text(s, 0.9, 3.5, 11.5, 1.5, [
        "The results slides from the previous version, unchanged.",
        "Every number from the talk is here with its full measurement detail, for questions.",
    ], size=18, color=INK2)
    notes(s, "Backup. The detailed results slides from v18, kept whole for "
             "questions.")
    return out


def rewrite_text_slide(prs, slide, ttl, sub, items):
    """Replace a v18 left-column slide's text box, keeping its picture."""
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip():
            sh._element.getparent().remove(sh._element)
    title(slide, ttl, sub, w=5.3)
    bullets(slide, 0.55, 1.7, 5.3, 4.5, items)


def main() -> int:
    shutil.copyfile(SRC, DST)
    prs = Presentation(DST)
    v18 = {i + 1: s for i, s in enumerate(prs.slides)}
    ids = {i + 1: el for i, el in enumerate(list(prs.slides._sldIdLst))}
    n0 = len(ids)
    assert n0 == 80, n0

    # light rewrites of two map-heavy 613590 slides that stay in the main flow
    rewrite_text_slide(prs, v18[69], "The road network it drew",
                       "613590, a tile no road model ever trained on.", [
        "231.4 km of road, in 3,693 segments.",
        "Raw output. Nothing corrected by hand.",
        "Cut-off 0.30. Each line traces the model's output. Nothing is invented to join gaps.",
    ])
    rewrite_text_slide(prs, v18[70], "Its roads against the public map",
                       "231 km found. The public map has 40 km.", [
        "TIGER is the US Census road layer.",
        "TIGER lists roads people drive on today.",
        "These are access tracks cut a century ago and left to grow over.",
        "That is the road network that leads to a forgotten well.",
    ])

    m = map_slides(prs)
    r = results_slides(prs, v18)

    def sid(slide):
        target = slide.slide_id
        for el in prs.slides._sldIdLst:
            if int(el.get("id")) == target:
                return el
        raise KeyError(target)

    order = []
    order += [ids[i] for i in range(1, 34)]                 # 1-33 as v18
    order += [sid(s) for s in m]                            # map section
    order += [ids[i] for i in range(34, 56)]                # 34-55 as v18
    order += [sid(r["score"]), sid(r["pits"]), sid(r["pads"]), sid(r["roads"]),
              ids[58], sid(r["summary"]), sid(r["arch"]),   # drainage kept
              ids[62], sid(r["tile2"]), ids[69], ids[70],
              sid(r["numbers"])]
    order += [ids[i] for i in range(73, 81)]                # limits .. appendix
    order += [sid(r["backup"])]
    order += [ids[i] for i in (56, 57, 59, 60, 61, 63, 64, 65, 66, 67, 68, 71, 72)]

    lst = prs.slides._sldIdLst
    assert len(order) == len(lst) == len(set(map(id, order))), (len(order), len(lst))
    for el in list(lst):
        lst.remove(el)
    for el in order:
        lst.append(el)

    prs.save(DST)
    print(f"wrote {DST}  ({len(order)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

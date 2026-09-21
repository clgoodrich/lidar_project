"""v10 -> v11: stop calling 613590 "held out", and move the locator to where it lands.

TERMINOLOGY: "HELD-OUT" WAS DOING TWO DIFFERENT JOBS
---------------------------------------------------
The deck used the same phrase for two unrelated things:

  slide 48   "the held-out blocks"   blocks inside the TRAINING tile that were
                                     withheld from training and kept for
                                     testing. This is the standard meaning and
                                     it is correct. It stays.

  9 slides   "Held-Out Tile 613590"  a different 4.5 km tile, 1.5 km west, that
                                     was never in any split at all.

The second use is wrong. "Held out" means withheld from a dataset the model
otherwise saw. 613590 was never part of that dataset, so there was nothing to
withhold. Using one label for both makes them look like the same kind of
evidence, when the second is the stronger claim.

WHY NOT "OUT-OF-DOMAIN"
-----------------------
It is the term the repo uses elsewhere (`LEADERBOARD.md`, `FIGURE_INDEX.md`)
and it is tempting, but on this deck it over-claims. 613590 is the same county,
the same 2019 survey, the same flight block and the same terrain as 9t -- 1.5
kilometres away. That is not a different domain. The conclusion slide already
says so: "Two tiles in one county is not evidence of regional generalisation."
Calling it out-of-domain here would contradict that.

So: **"Second tile, 613590"**. It is literally true, it claims nothing beyond
what was done, and it cannot be confused with the held-out blocks.

THE LOCATOR SLIDE
-----------------
Slide 7, "Where the two study tiles are", carried the line "9t is where every
model was trained. 613590 is 1.5 km west, and no model has ever seen it." At
slide 7 the audience has not yet met training, testing or splits, so the claim
lands on nothing. Moved to sit immediately before the second-tile results,
where it is the setup for them, and its subtitle rewritten to do that job.

Slide 6 "Study Area" keeps its own maps, so the early section is not left
without geographic context.

Run:
    python docs/presentation/_build_v11_fix_heldout_and_locator.py
Reads:  docs/presentation/WellSight_Presentation v10.pptx
Writes: docs/presentation/WellSight_Presentation v11.pptx
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs/presentation/WellSight_Presentation v10.pptx"
DST = ROOT / "docs/presentation/WellSight_Presentation v11.pptx"

LOCATOR = "Where the two study tiles are"
#: the first of the 613590 result slides; the locator goes immediately before it
FIRST_SECOND_TILE = "Pits"

OLD = re.compile(r"[Hh]eld[- ][Oo]ut\s+[Tt]ile\s+613590", re.U)
NEW = "Second tile, 613590"

LOCATOR_SUB = ("Everything so far was trained and tested on 9t. This is a "
               "different tile, 1.5 km west, that no model has ever seen.")
LOCATOR_NOTE = (
    "Before the next set of results, here is where they come from.\n\n"
    "Everything up to this point was trained and tested on the tile called 9t "
    "— trained on most of it, scored on blocks inside it that the model "
    "never saw.\n\n"
    "This is a different tile. 613590, a kilometre and a half west. No model "
    "has ever seen any part of it.\n\n"
    "Be careful with the word for this. It is not a 'held-out' tile, because "
    "it was never in the dataset to be held out of. It is simply a second "
    "tile.\n\n"
    "And it is not a generalisation test either. Same county, same 2019 "
    "survey, same flight block, same terrain. A tile from a different "
    "acquisition would be the real test, and that is on the further-research "
    "slide.")


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def retitle(slide) -> bool:
    """Swap the phrase wherever it appears in this slide's text and notes."""
    hit = False
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if OLD.search(r.text):
                    r.text = OLD.sub(NEW, r.text)
                    hit = True
    if slide.has_notes_slide:
        tf = slide.notes_slide.notes_text_frame
        if OLD.search(tf.text):
            tf.text = OLD.sub(NEW, tf.text)
            hit = True
    return hit


def main() -> int:
    prs = Presentation(SRC)
    lst = prs.slides._sldIdLst

    # ---- 1. rename, everywhere the phrase appears ----------------------
    renamed = []
    for i, s in enumerate(prs.slides, 1):
        if retitle(s):
            renamed.append((i, title_of(s)))
    print(f"  renamed on {len(renamed)} slide(s):")
    for i, t in renamed:
        print(f"    {i:3d}  {t}")

    # ---- 2. move the locator to sit before the second-tile results -----
    loc_i = next(i for i, s in enumerate(prs.slides)
                 if title_of(s).strip() == LOCATOR)
    tgt_i = next(i for i, s in enumerate(prs.slides)
                 if title_of(s).startswith(NEW))

    loc = prs.slides[loc_i]
    # reword the standfirst: it now introduces the section rather than the map
    texts = [sh for sh in loc.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    title_sh = min(texts, key=lambda s: s.top if s.top is not None else 0)
    for sh in texts:
        if sh is title_sh:
            continue
        runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
        if runs:
            runs[0].text = LOCATOR_SUB
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            break
    loc.notes_slide.notes_text_frame.text = LOCATOR_NOTE

    node = list(lst)[loc_i]
    lst.remove(node)
    # after removing the earlier slide everything below shifts up one
    ids = list(lst)
    ids[tgt_i - 1].addprevious(node)
    print(f"\n  moved {LOCATOR!r} from slide {loc_i + 1} to slide {tgt_i}")

    prs.save(DST)

    # ---- prove it -------------------------------------------------------
    out = Presentation(DST)
    left = []
    for i, s in enumerate(out.slides, 1):
        body = " ".join(sh.text_frame.text for sh in s.shapes
                        if sh.has_text_frame)
        note = (s.notes_slide.notes_text_frame.text
                if s.has_notes_slide else "")
        if OLD.search(body) or OLD.search(note):
            left.append(i)
    blocks = [i for i, s in enumerate(out.slides, 1)
              if "held-out blocks" in " ".join(
                  sh.text_frame.text for sh in s.shapes if sh.has_text_frame)]
    print(f"\n  'held-out tile 613590' remaining: {len(left)} {left}")
    print(f"  'held-out blocks' kept (correct usage) on slides: {blocks}")
    print(f"  {len(out.slides)} slides")
    for i, s in enumerate(out.slides, 1):
        if 5 <= i <= 8 or 55 <= i <= 62:
            print(f"    {i:3d}  {title_of(s)[:60]}")
    print(f"  {DST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

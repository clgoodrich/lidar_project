"""Add the NISAR further-research slide, sized to what NISAR can actually do.

THE HONEST FRAMING
------------------
The natural pitch is "use NISAR to find pits". It does not survive contact with
the numbers already measured over this tile in
docs/nisar_lidar_supplement_proposal.md:

    GCOV backscatter   10 m pixel     ->  100 m2
    GUNW InSAR         80 m pixel     -> 6400 m2
    our interior pits  mean 32.1 m2   ->  ~6 m across

A pit is about a third of one backscatter pixel and a two-hundredth of one
InSAR pixel. NISAR will not detect one, and the slide says so in its first
bullet rather than letting an audience member work it out and distrust the
rest.

What survives the scale check is better than detection anyway. Once lidar has
located candidates, the 80 m InSAR cell is the right size to ask whether the
ground over a cluster of them is *moving* -- which is the wellbore-integrity
question, and the one thing a single 2019 lidar flight can never answer. The
recon measured that this is feasible here: a snow-free fall pair over 9t held
coherence 0.50 with 94% of pixels above 0.3, against 0.14 for a mid-winter
pair. Seasonal planning, not sensor capability, is the constraint.

The caveat bullet is not optional. Everything measured so far is BETA
pre-calibration data and rests on a single pair, which is exactly the caution
already recorded in docs/iterations/BACKLOG.md ("NISAR language overshoots").
Validated CONUS products arrive ~July 2026.

Placed after "What would move it next" and before References: it is a research
direction, not one of the near-term actions on that slide.

Run:
    python docs/presentation/_add_nisar_further_research_slide_v6.py
Writes (in place):
    docs/presentation/WellSight_Presentation v6.pptx
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs/presentation/WellSight_Presentation v6.pptx"
FIG = ROOT / "docs/presentation/figures_30to45min/v6/nisar_pixel_vs_pit_scale_9t.png"

PML = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

# same constants the conclusion slides use, so this one does not look bolted on
INK = RGBColor(0x14, 0x1A, 0x1F)
INK2 = RGBColor(0x54, 0x5C, 0x63)
ACCENT = RGBColor(0x00, 0x96, 0xC7)

AFTER_TITLE = "What would move it next"
TITLE = "Further research — watching what we found"
KICKER = "NISAR: L-band radar, free, over the same ground every 12 days"
BULLETS = [
    "It cannot find a pit — 10 m backscatter, 80 m InSAR, against pits "
    "averaging 32 m²",
    "It can measure whether the ground over them is moving, to the millimetre",
    "Already tested here: a snow-free fall pair over 9t held coherence 0.50, "
    "94% of pixels usable",
    "Mid-winter failed at 0.14 — snow and freeze-thaw, so pairs must be "
    "leaf-off but snow-free",
    "That turns a one-time 2019 snapshot into a monitored surface, which is "
    "the wellbore-integrity question",
    "Caveat: beta pre-calibration data, one pair. Validated products from "
    "~July 2026",
]
NOTES = (
    "The obvious pitch is 'use NISAR to find pits'. It does not work and the "
    "slide leads with why: a NISAR backscatter pixel is 10 m, an InSAR pixel "
    "is 80 m, and our interior pits average 32 square metres, roughly six "
    "metres across. The figure shows two real annotated pits inside one 80 m "
    "InSAR cell.\n\n"
    "The useful question is the other direction. Lidar finds the candidates; "
    "NISAR then asks whether the ground over them is subsiding, every twelve "
    "days, for free. That is wellbore integrity, and a single 2019 flight "
    "cannot answer it at all.\n\n"
    "Feasibility is measured, not assumed: a snow-free fall pair over this "
    "tile held coherence 0.50 with 94 percent of pixels above 0.3. A "
    "mid-winter pair came back at 0.14, decorrelated by snow and freeze-thaw. "
    "So the constraint is acquisition season, not the sensor.\n\n"
    "Be straight about maturity if asked: everything so far is beta "
    "pre-calibration data and rests on one pair. Validated CONUS products "
    "start around July 2026, and nothing quantitative should be claimed "
    "before then."
)


def rgb_runs(tf, runs):
    first = True
    for text, size, bold, colour, space in runs:
        if first and len(tf.paragraphs) and not tf.paragraphs[0].runs:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT     # master centres these otherwise
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = colour
        r.font.name = "Calibri"


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def main() -> int:
    if not FIG.exists():
        raise SystemExit(f"figure missing: {FIG}\n"
                         "run _nisar_vs_lidar_pit_scale_9t.py first")
    prs = Presentation(DECK)

    src_i = next(i for i, s in enumerate(prs.slides)
                 if title_of(s) == AFTER_TITLE)
    src = prs.slides[src_i]

    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    # carry the background across, or the slide reverts to the master's
    bg = src._element.find(f"{PML}cSld").find(f"{PML}bg")
    if bg is not None:
        new._element.find(f"{PML}cSld").insert(0, copy.deepcopy(bg))

    tb = new.shapes.add_textbox(Inches(1.00), Inches(0.62),
                                Inches(11.5), Inches(0.9))
    # add_textbox defaults to wrap="none" + spAutoFit, which makes renderers
    # shrink the box around the text and re-centre it -- the title came out
    # centred while the body, which sets word_wrap, came out left.
    tb.text_frame.word_wrap = True
    rgb_runs(tb.text_frame, [(TITLE, 30, True, ACCENT, 0)])

    body = new.shapes.add_textbox(Inches(1.00), Inches(1.72),
                                  Inches(6.25), Inches(4.75))
    body.text_frame.word_wrap = True
    runs = [(KICKER, 17, True, INK, 13)]
    for b in BULLETS:
        runs.append(("•  " + b, 14, False, INK2, 9))
    rgb_runs(body.text_frame, runs)

    # figure is 10.6 x 9.4 in, so height is width x 0.887
    w = 5.55
    new.shapes.add_picture(str(FIG), Inches(7.55), Inches(1.60),
                           width=Inches(w), height=Inches(w * 0.887))

    foot = new.shapes.add_textbox(Inches(1.00), Inches(6.62),
                                  Inches(11.0), Inches(0.5))
    foot.text_frame.word_wrap = True
    rgb_runs(foot.text_frame,
             [("Colton Goodrich  ·  University of Houston",
               13, False, INK2, 0)])

    new.notes_slide.notes_text_frame.text = NOTES

    # add_slide appends; move it to sit right after "What would move it next"
    sld_lst = prs.slides._sldIdLst
    node = list(sld_lst)[-1]
    sld_lst.remove(node)
    list(sld_lst)[src_i].addnext(node)

    prs.save(DECK)
    out = Presentation(DECK)
    print(f"  added after slide {src_i + 1} ({AFTER_TITLE!r})")
    print(f"  now slide {src_i + 2}: {title_of(out.slides[src_i + 1])!r}")
    print(f"  {len(out.slides)} slides")
    print(f"  {DECK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

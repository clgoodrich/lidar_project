"""Retire the template-matching pipeline slides from v6, keeping what still holds.

WHY
---
Slides 60-65 are the last of the pre-U-Net deck. Four of them describe a method
that is not in the project any more, and two of them say things the conclusion
slides already say better:

  60  Step 2: Manual Annotation      redundant -- slides 30-35 cover annotation
                                     with the current counts; "861 pits" is a
                                     number from before the rim/floor split.
  61  Step 3: Template Matching      gone entirely. Normalised cross-correlation
                                     against a mean 17x17 m pit template is not
                                     how anything is detected now.
  62  Step 4: Classification         the XGBoost/LightGBM/HistGB ensemble,
                                     ROC-AUC 0.905, PR-AUC 0.212. Superseded by
                                     the held-out U-Net numbers on slides 47-59.
  63  Annotation Quality Control     KEPT. The 856 measured pits it checked are
                                     the same annotations the U-Nets train on,
                                     so the QC still applies. Moved next to the
                                     other annotation-preparation slides.
  64  Challenges & Limitations       REPURPOSED. Its bullet list duplicates
                                     slide 67 ("What it does not say"), but its
                                     figure -- DEP-listed wells with no terrain
                                     signature -- is the evidence for why we
                                     annotate at all. That argument lived on
                                     slide 60 with no picture, and the picture
                                     lived here with no argument. Joined up.
  65  Future Work                    duplicates slide 68, and two of its five
                                     bullets ("U-Net semantic segmentation",
                                     "geomorphon +31% PR-AUC") describe work
                                     that is either done or belongs to the dead
                                     pipeline. Leaving it would contradict the
                                     rest of the deck.

The "n / 4" step markers go with 61 and 62. Nothing in the current method is
numbered that way, so the marker is dropped rather than renumbered to a
sequence with no other members.

Run:
    python docs/presentation/_retire_template_matching_slides_v6.py
Writes (in place):
    docs/presentation/WellSight_Presentation v6.pptx
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs/presentation/WellSight_Presentation v6.pptx"

RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

#: 1-based, as the deck reads before any edit.
DROP = (60, 61, 62, 65)
REPURPOSE = 64
#: Annotation QC belongs beside the other label-preparation slides, not after
#: the outcomes. 38 is "Preprocessing - Annotations - Standard Values".
QC_SLIDE, QC_AFTER = 63, 38

NEW_TITLE = "Why we drew our own labels"
NEW_BULLETS = [
    "DEP records and visible terrain do not line up",
    "Many DEP dots have no pit on the ground beneath them",
    "Many clear pits carry no DEP record at all",
    "So DEP locations cannot serve as the training labels",
    "One derivative is not enough either — LRM and TPI alone overlap natural depressions",
]
#: Slide 63's closing line still pointed at the retired classifier.
QC_OLD_TAIL = "before training the classifier on them"
QC_NEW_TAIL = "before any model is trained on them"


def set_lines(tf, lines):
    """Rewrite a text frame's paragraphs, keeping each one's first run's font.

    Deleting paragraphs and adding fresh ones loses the theme run properties and
    the bullet comes back as plain body text, so reuse the paragraphs that are
    already there and clone the first one when more are needed.
    """
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):
        new = copy.deepcopy(paras[0]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for p, text in zip(paras, lines):
        runs = list(p.runs)
        if not runs:                       # empty paragraph, nothing to inherit
            p.text = text
            continue
        runs[0].text = text
        for r in runs[1:]:                 # drop the tail, keep run 0's font
            r._r.getparent().remove(r._r)
    for p in paras[len(lines):]:           # surplus paragraphs go
        p._p.getparent().remove(p._p)


def main() -> int:
    prs = Presentation(DECK)
    sld_lst = prs.slides._sldIdLst
    before = len(prs.slides)
    # grab the doomed slides as objects NOW: the QC move below shifts every
    # index after 38, so DROP's 1-based numbers stop being valid the moment it
    # happens.
    doomed = [prs.slides[i - 1] for i in DROP]

    # -- repurpose 64 -----------------------------------------------------
    s = prs.slides[REPURPOSE - 1]
    texts = [sh for sh in s.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    # title is the topmost; the bullet block is the tall one on the left
    title = min(texts, key=lambda sh: sh.top)
    bullets = max((sh for sh in texts if sh is not title),
                  key=lambda sh: sh.height)
    set_lines(title.text_frame, [NEW_TITLE])
    set_lines(bullets.text_frame, NEW_BULLETS)
    print(f"  slide {REPURPOSE}: -> {NEW_TITLE!r}, {len(NEW_BULLETS)} bullets")

    # -- fix the QC slide's stale closing line ----------------------------
    for sh in prs.slides[QC_SLIDE - 1].shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if QC_OLD_TAIL in r.text:
                    r.text = r.text.replace(QC_OLD_TAIL, QC_NEW_TAIL)
                    print(f"  slide {QC_SLIDE}: closing line updated")

    # -- move the QC slide up, before anything is deleted -----------------
    ids = list(sld_lst)
    node = ids[QC_SLIDE - 1]
    sld_lst.remove(node)
    ids = list(sld_lst)
    ids[QC_AFTER - 1].addnext(node)
    print(f"  slide {QC_SLIDE}: moved to sit after slide {QC_AFTER}")

    # -- delete, by identity against the objects grabbed before the move ---
    for slide in doomed:
        rid = None
        for sid in list(sld_lst):
            if prs.part.rels[sid.get(RID)].target_part is slide.part:
                rid = sid.get(RID)
                sld_lst.remove(sid)
                break
        if rid:
            prs.part.drop_rel(rid)
    print(f"  dropped slides {DROP}")

    prs.save(DECK)
    print(f"\n  {before} slides -> {len(Presentation(DECK).slides)}")
    print(f"  {DECK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

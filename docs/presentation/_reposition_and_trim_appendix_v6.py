"""Second pass: place the salvaged slide where it argues, drop the orphan appendix.

Two things the first pass left.

1. "Why we drew our own labels" landed at 61, after every held-out result,
   because it inherited slide 64's position. It is the motivation for
   annotating at all, so it belongs immediately before the annotation section
   (slide 30, "Manual Annotation - Roads") rather than among the conclusions.

2. "Appendix: 48-Feature Set (1/2)" and "(2/2)" tabulate the per-candidate
   features of the XGBoost/LightGBM ensemble. That model came out of the deck
   in the first pass, so the tables now document something the audience is
   never shown. Dropped.

   NOT dropped: "Appendix: Literature-Grounded Parameters". Its TPI radii, LRM
   windows, roughness kernel and openness radius are the values the current
   derivative builder still uses, with the citations behind them. Two rows in
   it are worth a look before the talk, but they are not wrong enough to change
   silently:
     - "DEM resolution 1 m" while the delivered derivative stack is 0.5 m
     - "Blob sigma range 0.8-5.0", which was a parameter of the retired
       blob-detection stage
   Both left as they are, and reported to the user instead.

Run:
    python docs/presentation/_reposition_and_trim_appendix_v6.py
Writes (in place):
    docs/presentation/WellSight_Presentation v6.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs/presentation/WellSight_Presentation v6.pptx"
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

MOVE_TITLE = "Why we drew our own labels"
#: sit immediately before this one, so the argument precedes the annotation run
MOVE_BEFORE_TITLE = "Manual Annotation – Roads"
DROP_TITLES = ("Appendix: 48-Feature Set (1/2)",
               "Appendix: 48-Feature Set (2/2)")


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


def find(prs, title):
    for i, s in enumerate(prs.slides):
        t = title_of(s)
        # the deck mixes hyphen and en dash in titles, so compare loosely
        if t.replace("–", "-").replace("—", "-") == \
           title.replace("–", "-").replace("—", "-"):
            return i, s
    raise SystemExit(f"could not find slide titled {title!r}")


def main() -> int:
    prs = Presentation(DECK)
    sld_lst = prs.slides._sldIdLst
    before = len(prs.slides)

    doomed = [find(prs, t)[1] for t in DROP_TITLES]
    _, mover = find(prs, MOVE_TITLE)

    # -- move first, while every slide is still present -------------------
    ids = list(sld_lst)
    node = ids[[s.part for s in prs.slides].index(mover.part)]
    sld_lst.remove(node)
    tgt_i, _ = find(prs, MOVE_BEFORE_TITLE)
    list(sld_lst)[tgt_i].addprevious(node)
    print(f"  {MOVE_TITLE!r} -> just before {MOVE_BEFORE_TITLE!r}")

    # -- then delete ------------------------------------------------------
    for slide in doomed:
        for sid in list(sld_lst):
            if prs.part.rels[sid.get(RID)].target_part is slide.part:
                sld_lst.remove(sid)
                prs.part.drop_rel(sid.get(RID))
                break
    print(f"  dropped {len(doomed)} orphaned appendix slides")

    prs.save(DECK)
    print(f"\n  {before} slides -> {len(Presentation(DECK).slides)}")
    print(f"  {DECK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

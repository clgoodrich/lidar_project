from pptx import Presentation
from pptx.util import Inches

PPTX = 'WellSight_Presentation.pptx'
prs = Presentation(PPTX)

# Current slide order:
# 6: DEP Records Problem
# 9: Step 2: Manual Annotation
# 14: Challenges & Limitations
# 15: Future Work

# Replace graphics on slides 6, 9, 14 with new versions
# Restore slide 15 (Future Work) to original image

replacements = {
    6:  'close up errors.png',
    9:  'Manual Picks.png',
    14: 'orphaned with no clear wells.png',
    15: '_future_work_img.png',  # the original we extracted earlier
}

for si, img in replacements.items():
    slide = prs.slides[si]
    # Find existing picture position/size before removing
    old_left, old_top, old_w, old_h = None, None, None, None
    for shape in list(slide.shapes):
        if shape.shape_type == 13:
            old_left = shape.left
            old_top = shape.top
            old_w = shape.width
            old_h = shape.height
            shape._element.getparent().remove(shape._element)

    if old_left is not None:
        slide.shapes.add_picture(img, old_left, old_top, old_w, old_h)
    else:
        slide.shapes.add_picture(img, Inches(5.0), Inches(1.3), Inches(4.8), Inches(5.8))

    title = ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            title = shape.text_frame.text[:40]
            break
    print(f"Replaced graphic on slide {si} ({title})")

# Update Step 4 notes with ROC-AUC clarity
for i, slide in enumerate(prs.slides):
    for shape in slide.shapes:
        if shape.has_text_frame and 'Classification' in shape.text_frame.text:
            ns = slide.notes_slide
            ns.notes_text_frame.text = """STEP 4: CLASSIFICATION & RESULTS

WHAT WE DID
- Started with 628,332 candidate locations from template matching
- Measured 48 terrain properties at each one (depth, shape, roughness, symmetry, etc.)
- Trained 3 machine learning models to learn "real pit" vs "not a pit"
- Combined their votes into one final probability score per candidate

THE PROBABILITY SCORE
- Each candidate gets a score from 0 to 1
- This comes from averaging the 3 models' outputs, then calibrating so the number reflects reality
- If the model says 0.80, it means roughly 80% of candidates scored 0.80 turned out to be real pits
- The score is the model's confidence that a given spot is a real well pit

THE THRESHOLD — CHOOSING A CUTOFF
- You pick a minimum score and throw away everything below it
- At 0.70: keep 507 candidates, 76.9% are real pits (wide net, more false positives)
- At 0.80: keep 289 candidates, 85.5% are real (good balance)
- At 0.90: keep 118 candidates, 93.2% are real (tight filter, high confidence)
- It is a dial: wide net for survey planning vs. tight filter for expensive field visits

ROC-AUC = 0.905 — THE MODEL'S REPORT CARD
- This measures overall discrimination ability across ALL thresholds, not just one
- Literal meaning: pick one real pit and one random non-pit. 90.5% of the time, the model gives the real pit a higher score.
- It answers: "Did the model learn what a pit looks like?" Yes — 0.905 is very good.
- 0.5 = coin flip (useless). 1.0 = perfect. 0.905 = strong.
- ROC-AUC does NOT tell you how many candidates to keep — that is the threshold's job.

PR-AUC = 0.212 — THE HONEST CHECK
- Precision-Recall AUC accounts for how rare pits are (only 1 in 270 candidates is real)
- A random guesser would get PR-AUC = 0.0037. Ours is 0.212 — 57x better than random.
- PR-AUC is always lower than ROC-AUC when positives are rare. This is normal.
- It answers: "Can the model find needles in a haystack?" Yes, but it is hard.

THE THREE MODELS
- XGBoost, LightGBM, HistGradientBoosting — all "gradient boosted decision trees"
- Each makes predictions through hundreds of chained yes/no questions about terrain properties
- Using 3 instead of 1 reduces blind spots from any single model
- Isotonic calibration converts raw scores to real probabilities

SPATIAL CROSS-VALIDATION
- The map is split into zones with 500m gaps between training and testing areas
- Without gaps, nearby terrain is so similar the model would cheat and score artificially high
- 5-fold: each area takes a turn as the test set
- This gives honest, real-world performance numbers

THE GRAPHIC
- Green circles = pits we annotated by hand (ground truth)
- Colored dots = pits the model found (yellow = ~0.6 confidence, red = ~1.0)
- Where they overlap = correct detection
"""
            print(f"Updated notes for slide {i} (Classification)")
            break

prs.save(PPTX)
print("\nDone.")

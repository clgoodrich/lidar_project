from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from lxml import etree
import copy
import os

PPTX = 'WellSight_Presentation.pptx'
TEMPLATE_IMG = r'C:\Users\colto\Documents\GitHub\lidar_project\data\derivatives\pit_1m_template_mean.png'
WELLSPOT_IMG = 'wellspot.jpg'
PROFILE_IMG = 'pit_profile_diagram.png'

prs = Presentation(PPTX)
slides = prs.slides
slide_width = prs.slide_width
slide_height = prs.slide_height

# --- TASK 1: Replace graphic on slide 7 (Template Matching) ---
slide7 = slides[7]
# Remove existing picture shapes
to_remove = []
for shape in slide7.shapes:
    if shape.shape_type == 13:  # Picture
        to_remove.append(shape)
for shape in to_remove:
    sp = shape._element
    sp.getparent().remove(sp)
    print(f"  Removed picture from slide 7")

# Add new template image - fill most of slide below title
pic = slide7.shapes.add_picture(TEMPLATE_IMG, Inches(0.3), Inches(1.4), Inches(9.4), Inches(5.8))
print("Inserted new template graphic on slide 7")

# --- TASK 2: Split slide 9 (What a Well Pit Looks Like) ---
# Current slide 9 has both wellspot photo and cross-section
# Plan: 
#   - Add a new slide after slide 4 (Pipeline Overview) showing just the wellspot photo
#     Title: "What Does a Well Pit Look Like?"
#   - Keep existing slide 9 (which will become slide 10 after insertion) but 
#     make it just the cross-section. Title: "Measured Pit Morphology"

# First, let's see what's on slide 9
slide9 = slides[9]
print("\nSlide 9 shapes:")
for shape in slide9.shapes:
    if shape.shape_type == 13:
        print(f"  Picture: {shape.left}, {shape.top}, {shape.width}x{shape.height}")
    elif shape.has_text_frame:
        print(f"  Text: '{shape.text_frame.text[:50]}'")

# Remove all pictures from slide 9 - we'll add just the cross-section
pics_on_9 = []
for shape in slide9.shapes:
    if shape.shape_type == 13:
        pics_on_9.append(shape)
for shape in pics_on_9:
    sp = shape._element
    sp.getparent().remove(sp)
print(f"  Removed {len(pics_on_9)} pictures from slide 9")

# Change slide 9 title to "Measured Pit Morphology"
for shape in slide9.shapes:
    if shape.has_text_frame:
        tf = shape.text_frame
        if 'Well Pit' in tf.text or 'well pit' in tf.text.lower():
            for para in tf.paragraphs:
                for run in para.runs:
                    run.text = ""
            tf.paragraphs[0].runs[0].text = "Measured Pit Morphology"
            print("  Updated slide 9 title to 'Measured Pit Morphology'")
            break

# Add cross-section image to slide 9
pic = slide9.shapes.add_picture(PROFILE_IMG, Inches(0.5), Inches(1.4), Inches(9.0), Inches(5.8))
print("  Added cross-section to slide 9")

# --- Now insert a new slide after slide 4 for the wellspot photo ---
# python-pptx doesn't have a direct "insert slide at position" method,
# but we can add a slide and move it via XML manipulation

# Add new slide using blank layout
blank_layout = None
for layout in prs.slide_layouts:
    if 'blank' in layout.name.lower() or layout.name == 'Blank':
        blank_layout = layout
        break
if blank_layout is None:
    blank_layout = prs.slide_layouts[6]  # typically blank

new_slide = prs.slides.add_slide(blank_layout)

# Set background to dark
bg = new_slide.background
fill = bg.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

# Add title text box
from pptx.util import Inches, Pt
txBox = new_slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.9))
tf = txBox.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "What Does a Well Pit Look Like?"
run.font.size = Pt(28)
run.font.bold = True
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Add wellspot photo - centered, large
pic = new_slide.shapes.add_picture(WELLSPOT_IMG, Inches(1.5), Inches(1.3), Inches(7.0), Inches(5.5))

# Add caption
capBox = new_slide.shapes.add_textbox(Inches(1.5), Inches(6.9), Inches(7.0), Inches(0.5))
tf2 = capBox.text_frame
p2 = tf2.paragraphs[0]
p2.alignment = PP_ALIGN.CENTER
run2 = p2.add_run()
run2.text = "Exposed wellhead in overgrown brush (Scott Detrow / StateImpact PA, 2012)"
run2.font.size = Pt(11)
run2.font.italic = True
run2.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

print("\nAdded new wellspot slide (currently at end)")

# Now move the new slide (last) to position after slide 4 (index 5)
# Access the slide list XML
sldIdLst = prs.slides._sldIdLst
sldId_elements = list(sldIdLst)
# The new slide is the last one
new_sldId = sldId_elements[-1]
# Remove it from current position
sldIdLst.remove(new_sldId)
# Insert after slide 4 (index 4), so at position 5
sldIdLst.insert(5, new_sldId)
print("Moved wellspot slide to position 5 (after Pipeline Overview)")

# Save
prs.save(PPTX)
print(f"\nSaved {PPTX}")

# Print final slide order
prs2 = Presentation(PPTX)
for i, slide in enumerate(prs2.slides):
    title = ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                title = txt[:60]
                break
    print(f"Slide {i}: {title}")

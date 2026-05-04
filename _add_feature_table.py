from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PPTX = 'WellSight_Presentation.pptx'
prs = Presentation(PPTX)

BG = RGBColor(0x1a, 0x1a, 0x2e)
ACCENT = RGBColor(0x00, 0x96, 0xC7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xCC, 0xCC, 0xCC)
ROW_DARK = RGBColor(0x16, 0x16, 0x2a)
ROW_LIGHT = RGBColor(0x22, 0x22, 0x3a)
HEADER_BG = RGBColor(0x00, 0x64, 0x8C)

features = [
    (1, "lrm_5_inner_min", "Depth/contrast", "Min LRM-5 in inner ring"),
    (2, "lrm_5_inner_mean", "Depth/contrast", "Mean LRM-5 in inner ring"),
    (3, "lrm_5_rim_mean", "Depth/contrast", "Mean LRM-5 in rim ring"),
    (4, "lrm_5_rim_minus", "Depth/contrast", "Rim minus inner min (LRM-5)"),
    (5, "lrm_5_sym", "Depth/contrast", "Radial symmetry std (LRM-5)"),
    (6, "lrm_11_inner_min", "Depth/contrast", "Min LRM-11 in inner ring"),
    (7, "lrm_11_inner_mean", "Depth/contrast", "Mean LRM-11 in inner ring"),
    (8, "lrm_11_rim_mean", "Depth/contrast", "Mean LRM-11 in rim ring"),
    (9, "lrm_11_rim_minus", "Depth/contrast", "Rim minus inner min (LRM-11)"),
    (10, "lrm_11_sym", "Depth/contrast", "Radial symmetry std (LRM-11)"),
    (11, "lrm_25_inner_min", "Depth/contrast", "Min LRM-25 in inner ring"),
    (12, "lrm_25_inner_mean", "Depth/contrast", "Mean LRM-25 in inner ring"),
    (13, "lrm_25_rim_mean", "Depth/contrast", "Mean LRM-25 in rim ring"),
    (14, "lrm_25_rim_minus", "Depth/contrast", "Rim minus inner min (LRM-25)"),
    (15, "lrm_25_sym", "Depth/contrast", "Radial symmetry std (LRM-25)"),
    (16, "tpi_05_inner_min", "Depth/contrast", "Min TPI-5 in inner ring"),
    (17, "tpi_05_inner_mean", "Depth/contrast", "Mean TPI-5 in inner ring"),
    (18, "tpi_05_rim_mean", "Depth/contrast", "Mean TPI-5 in rim ring"),
    (19, "tpi_05_rim_minus", "Depth/contrast", "Rim minus inner min (TPI-5)"),
    (20, "tpi_05_sym", "Depth/contrast", "Radial symmetry std (TPI-5)"),
    (21, "tpi_15_inner_min", "Depth/contrast", "Min TPI-15 in inner ring"),
    (22, "tpi_15_inner_mean", "Depth/contrast", "Mean TPI-15 in inner ring"),
    (23, "tpi_15_rim_mean", "Depth/contrast", "Mean TPI-15 in rim ring"),
    (24, "tpi_15_rim_minus", "Depth/contrast", "Rim minus inner min (TPI-15)"),
    (25, "tpi_15_sym", "Depth/contrast", "Radial symmetry std (TPI-15)"),
    (26, "openness_neg_inner_min", "Depth/contrast", "Min neg. openness in inner ring"),
    (27, "openness_neg_inner_mean", "Depth/contrast", "Mean neg. openness in inner ring"),
    (28, "openness_neg_rim_mean", "Depth/contrast", "Mean neg. openness in rim ring"),
    (29, "openness_neg_rim_minus", "Depth/contrast", "Rim minus inner min (neg. openness)"),
    (30, "openness_neg_sym", "Depth/contrast", "Radial symmetry std (neg. openness)"),
    (31, "slope_inner_mean", "Surface", "Mean slope in inner ring"),
    (32, "slope_window_max", "Surface", "Max slope in 17x17 m window"),
    (33, "roughness_5_inner_mean", "Surface", "Mean roughness in inner ring"),
    (34, "roughness_5_window_max", "Surface", "Max roughness in window"),
    (35, "local_relief_10_inner_mean", "Surface", "Mean local relief (10 m) in inner ring"),
    (36, "local_relief_10_window_max", "Surface", "Max local relief (10 m) in window"),
    (37, "chm_inner_mean", "Surface", "Mean canopy height in inner ring"),
    (38, "chm_window_max", "Surface", "Max canopy height in window"),
    (39, "density_inner_mean", "Density/match", "Mean ground point density, inner ring"),
    (40, "density_window_mean", "Density/match", "Mean ground point density, window"),
    (41, "dem_cut_m", "Density/match", "DEM rim-minus-inner contrast (raw elev.)"),
    (42, "match_score_center", "Density/match", "NCC template match score at center"),
    (43, "morph_center_depth", "Morphology", "LRM-5 value at center pixel"),
    (44, "morph_ring_r0_2", "Morphology", "Mean LRM-5 in 0-2 m band"),
    (45, "morph_ring_r2_5", "Morphology", "Mean LRM-5 in 2-5 m band"),
    (46, "morph_ring_r5_8", "Morphology", "Mean LRM-5 in 5-8 m band"),
    (47, "morph_radial_monotonicity", "Morphology", "Spearman corr. of ring means vs distance"),
    (48, "morph_compactness", "Morphology", "Fraction of LRM signal in inner ring"),
]

def add_feature_slide(prs, title_text, subtitle_text, feature_subset):
    blank = None
    for layout in prs.slide_layouts:
        if 'blank' in layout.name.lower():
            blank = layout
            break
    if blank is None:
        blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)

    # Background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(9), Inches(0.5))
    p = txBox.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title_text
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = ACCENT

    # Subtitle
    txBox2 = slide.shapes.add_textbox(Inches(0.4), Inches(0.7), Inches(9), Inches(0.3))
    p2 = txBox2.text_frame.paragraphs[0]
    run2 = p2.add_run()
    run2.text = subtitle_text
    run2.font.size = Pt(13)
    run2.font.color.rgb = GRAY

    # Table
    n_rows = len(feature_subset) + 1  # +1 for header
    n_cols = 4
    left = Inches(0.3)
    top = Inches(1.1)
    width = Inches(9.4)
    height = Inches(6.1)

    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table

    # Column widths
    table.columns[0].width = Inches(0.4)   # #
    table.columns[1].width = Inches(2.8)   # Feature Name
    table.columns[2].width = Inches(1.5)   # Category
    table.columns[3].width = Inches(4.7)   # Description

    # Header row
    headers = ["#", "Feature Name", "Category", "Description"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for para in cell.text_frame.paragraphs:
            para.font.size = Pt(10)
            para.font.bold = True
            para.font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_BG

    # Data rows
    for i, (num, name, cat, desc) in enumerate(feature_subset):
        row_idx = i + 1
        row_color = ROW_DARK if i % 2 == 0 else ROW_LIGHT
        for j, val in enumerate([str(num), name, cat, desc]):
            cell = table.cell(row_idx, j)
            cell.text = val
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(9)
                para.font.color.rgb = WHITE if j < 2 else GRAY
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_color

    return slide

# Split: features 1-30 (Depth/contrast) on slide 1, features 31-48 (Surface/Density/Morphology) on slide 2
slide1 = add_feature_slide(prs,
    "Appendix: 48-Feature Set (1/2)",
    "Depth/contrast features: inner ring (0-3 m) vs rim ring (6-8 m) comparisons across terrain derivatives",
    features[:30])

slide2 = add_feature_slide(prs,
    "Appendix: 48-Feature Set (2/2)",
    "Surface, density/match, and morphology features",
    features[30:])

# Move both slides to end (they already are at end, so just verify)
prs.save(PPTX)

# Verify
prs2 = Presentation(PPTX)
for i, slide in enumerate(prs2.slides):
    title = ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                title = txt[:60]
                break
    if i >= 15:
        print(f"Slide {i}: {title}")

print(f"\nTotal slides: {len(prs2.slides)}")
print("Done.")
